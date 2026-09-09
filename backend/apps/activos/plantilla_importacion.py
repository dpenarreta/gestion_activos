"""Plantilla .xlsx para la carga masiva de activos.

Se genera al vuelo en vez de servir un archivo estático porque incluye los
catálogos vigentes —tipos de dispositivo, departamentos y empleados que
existen *hoy*— en hojas aparte. Un archivo fijo quedaría desactualizado en
cuanto alguien registre un área nueva, y quien llena la plantilla tendría que
adivinar qué códigos son válidos o descubrirlos por ensayo y error.

La hoja de captura queda **vacía bajo los encabezados**: las ayudas van como
comentarios de celda y el ejemplo en una hoja aparte. Ponerlos como filas
dentro de la hoja de datos obligaría a borrarlos antes de cargar, y olvidarlo
produciría filas basura o errores de validación desconcertantes.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from apps.organizacion.models import Departamento, Empleado, Proveedor, Sede

from .importacion import MAX_FILAS, NOMBRE_HOJA_DATOS, columnas_configuradas
from .models import TipoDispositivo

RELLENO_OBLIGATORIA = PatternFill("solid", fgColor="D9E2F3")
RELLENO_OPCIONAL = PatternFill("solid", fgColor="EDEDED")
RELLENO_CABECERA_AUX = PatternFill("solid", fgColor="F2F2F2")

ANCHOS = {
    "tipo": 22,
    "nombre": 30,
    "marca": 16,
    "modelo": 20,
    "numero_serie": 22,
    "departamento": 22,
    "fecha_adquisicion": 20,
    "custodio": 24,
    "sede": 26,
    "criticidad": 14,
    "uso": 18,
    "fecha_ingreso": 18,
    "costo_adquisicion": 16,
    "proveedor": 24,
    "fecha_fin_garantia": 18,
    "especificaciones": 44,
    "observaciones": 30,
}
ANCHO_POR_DEFECTO = 24

EJEMPLOS = [
    {
        "tipo": "LAP",
        "nombre": "Laptop Contabilidad 02",
        "marca": "Dell",
        "modelo": "Latitude 5440",
        "numero_serie": "DL5440-00002",
        "departamento": "CTB",
        "fecha_adquisicion": "2024-03-15",
        "custodio": "TI-0001",
        "sede": "Matriz Quito",
        "criticidad": "Alta",
        "uso": "Administrativo",
        "fecha_ingreso": "2024-03-20",
        "costo_adquisicion": "1150.00",
        "proveedor": "Tecnomega",
        "fecha_fin_garantia": "2027-03-15",
        "especificaciones": "Procesador=Intel i5-1335U; RAM=16 GB; Disco=512 GB SSD",
        "observaciones": "Incluye maletín y mouse",
    },
    {
        "tipo": "IMP",
        "nombre": "Impresora Recepción",
        "marca": "HP",
        "modelo": "LaserJet M404",
        "numero_serie": "HPM404-77120",
        "departamento": "CTB",
        "fecha_adquisicion": "2023-11-02",
        "custodio": "",
        "sede": "Sucursal Guayaquil",
        "criticidad": "Baja",
        "uso": "Atención al cliente",
        "fecha_ingreso": "",
        "costo_adquisicion": "320",
        "proveedor": "Comptronix",
        "fecha_fin_garantia": "",
        "especificaciones": "Resolución=1200 dpi; Conectividad=Ethernet",
        "observaciones": "Sin custodio: queda en bodega",
    },
]


def _escribir_encabezados(hoja, columnas, con_ayudas: bool) -> None:
    for indice, columna in enumerate(columnas, start=1):
        celda = hoja.cell(row=1, column=indice, value=columna.encabezado)
        celda.font = Font(bold=True)
        celda.fill = RELLENO_OBLIGATORIA if columna.obligatoria else RELLENO_OPCIONAL
        celda.alignment = Alignment(vertical="center", wrap_text=True)
        if con_ayudas and columna.ayuda:
            comentario = Comment(columna.ayuda, "Gestión de Activos")
            comentario.width = 320
            comentario.height = 110
            celda.comment = comentario
        hoja.column_dimensions[get_column_letter(indice)].width = ANCHOS.get(
            columna.clave, ANCHO_POR_DEFECTO
        )
    hoja.freeze_panes = "A2"


def _hoja_datos(libro: Workbook, columnas) -> None:
    """Hoja de captura: encabezados y nada más."""
    hoja = libro.active
    hoja.title = NOMBRE_HOJA_DATOS
    _escribir_encabezados(hoja, columnas, con_ayudas=True)

    # La fecha es el error de captura más frecuente; validarla en Excel lo
    # ataja antes de subir el archivo.
    validacion_fecha = DataValidation(
        type="date",
        operator="greaterThanOrEqual",
        formula1="1990-01-01",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Fecha inválida",
        error="Escriba la fecha como AAAA-MM-DD, por ejemplo 2024-03-15.",
    )
    claves = [columna.clave for columna in columnas]
    if "fecha_adquisicion" in claves:
        hoja.add_data_validation(validacion_fecha)
        columna_fecha = get_column_letter(claves.index("fecha_adquisicion") + 1)
        validacion_fecha.add(f"{columna_fecha}2:{columna_fecha}{MAX_FILAS + 1}")


def _hoja_ejemplo(libro: Workbook, columnas) -> None:
    """Mismas columnas, con filas de muestra. Separada de la de captura para
    que no haya nada que borrar antes de importar."""
    hoja = libro.create_sheet("Ejemplo")
    _escribir_encabezados(hoja, columnas, con_ayudas=False)
    for numero_fila, ejemplo in enumerate(EJEMPLOS, start=2):
        for indice, columna in enumerate(columnas, start=1):
            # Una columna propia no tiene ejemplo predefinido: se deja vacía
            # antes que inventar un valor que confunda.
            hoja.cell(row=numero_fila, column=indice, value=ejemplo.get(columna.clave, ""))


def _hoja_catalogo(libro: Workbook, titulo: str, encabezados: list[str], filas: list[list]) -> None:
    hoja = libro.create_sheet(titulo)
    for indice, encabezado in enumerate(encabezados, start=1):
        celda = hoja.cell(row=1, column=indice, value=encabezado)
        celda.font = Font(bold=True)
        celda.fill = RELLENO_CABECERA_AUX
        hoja.column_dimensions[get_column_letter(indice)].width = 30
    for numero_fila, fila in enumerate(filas, start=2):
        for indice, valor in enumerate(fila, start=1):
            hoja.cell(row=numero_fila, column=indice, value=valor)
    if not filas:
        hoja.cell(row=2, column=1, value="(no hay registros: créelos antes de importar)")
    hoja.freeze_panes = "A2"


def _hoja_instrucciones(libro: Workbook, columnas) -> None:
    hoja = libro.create_sheet("Instrucciones", 0)
    hoja.column_dimensions["A"].width = 108

    lineas = [
        ("Carga masiva de activos", True),
        ("", False),
        (
            f"1. Complete la hoja «{NOMBRE_HOJA_DATOS}», una fila por equipo, desde la fila 2.",
            False,
        ),
        ("2. Las columnas marcadas con * son obligatorias.", False),
        (
            "3. Pase el cursor sobre cada encabezado para ver qué se espera en esa columna. "
            "En la hoja «Ejemplo» hay dos filas ya completadas.",
            False,
        ),
        (
            "4. En «Tipo de dispositivo» y «Departamento» escriba el código o el nombre exacto: "
            "los valores válidos están en las hojas «Tipos» y «Departamentos».",
            False,
        ),
        (
            "5. En «Código del custodio» use el código interno del empleado (hoja «Empleados»). "
            "Si lo deja vacío, el activo queda en bodega, sin responsable asignado.",
            False,
        ),
        (
            "6. La sede debe existir en el catálogo: copie el nombre exacto de la hoja "
            "«Sedes». Si la deja vacía, el activo queda registrado sin sitio hasta que "
            "alguien lo ubique.",
            False,
        ),
        (
            "7. El código de barras NO se llena: lo genera el sistema al importar, con el "
            "formato GA-<TIPO>-<SECUENCIA>.",
            False,
        ),
        ("", False),
        ("Qué pasa al cargar el archivo", True),
        (
            "El sistema revisa el archivo completo y muestra un reporte con los errores fila por "
            "fila. Nada se guarda hasta que usted confirme.",
            False,
        ),
        (
            "La importación es todo o nada: si una fila falla, no se carga ninguna. Corrija lo "
            "señalado y vuelva a subir el archivo.",
            False,
        ),
        (f"Máximo {MAX_FILAS} filas por carga.", False),
        ("", False),
        ("Columnas de esta plantilla", True),
        *[
            (
                f"· {columna.encabezado}" + (f" — {columna.ayuda}" if columna.ayuda else ""),
                False,
            )
            for columna in columnas
        ],
        ("", False),
        ("Errores frecuentes", True),
        ("· Número de serie repetido, dentro del archivo o ya existente en el inventario.", False),
        ("· Fecha en otro formato: use AAAA-MM-DD (por ejemplo 2024-03-15).", False),
        ("· Código de área o de tipo que no existe, o que está desactivado.", False),
        ("· Costo con símbolo de moneda: escriba solo el número (1150.00, no $1.150,00).", False),
        (
            "· Especificaciones sin el signo «=»: el formato es Clave=valor, separando cada par "
            "con punto y coma.",
            False,
        ),
    ]

    for numero_fila, (texto, es_titulo) in enumerate(lineas, start=1):
        celda = hoja.cell(row=numero_fila, column=1, value=texto)
        celda.alignment = Alignment(wrap_text=True, vertical="top")
        if es_titulo:
            celda.font = Font(bold=True, size=12)


def construir_plantilla() -> bytes:
    """Genera el .xlsx con la hoja de captura, las instrucciones y los catálogos.

    Las columnas salen de la configuración vigente, de modo que la plantilla
    refleja siempre lo que el importador va a leer.
    """
    columnas = columnas_configuradas()
    libro = Workbook()
    _hoja_datos(libro, columnas)
    _hoja_ejemplo(libro, columnas)

    _hoja_catalogo(
        libro,
        "Tipos",
        ["Código", "Nombre"],
        [
            [tipo.codigo, tipo.nombre]
            for tipo in TipoDispositivo.objects.filter(activo=True).order_by("nombre")
        ],
    )
    _hoja_catalogo(
        libro,
        "Departamentos",
        ["Código", "Nombre"],
        [
            [departamento.codigo, departamento.nombre]
            for departamento in Departamento.objects.filter(activo=True).order_by("nombre")
        ],
    )
    _hoja_catalogo(
        libro,
        "Sedes",
        ["Nombre", "Ciudad"],
        [
            [sede.nombre, sede.ciudad]
            for sede in Sede.objects.filter(activa=True).order_by("nombre")
        ],
    )
    _hoja_catalogo(
        libro,
        "Proveedores",
        ["Nombre", "Identificación"],
        [
            [proveedor.nombre, proveedor.identificacion]
            for proveedor in Proveedor.objects.filter(activo=True).order_by("nombre")
        ],
    )
    _hoja_catalogo(
        libro,
        "Empleados",
        ["Código", "Nombre", "Departamento"],
        [
            [empleado.codigo_empleado, empleado.nombre_completo, empleado.departamento.nombre]
            for empleado in Empleado.objects.filter(activo=True)
            .select_related("departamento")
            .order_by("apellidos", "nombres")
        ],
    )

    _hoja_instrucciones(libro, columnas)
    # Se abre en Instrucciones: es lo primero que conviene leer.
    libro.active = 0

    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()
