"""Dashboard principal (§15) y exportación a Excel (§16)."""

import datetime
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.activos.models import Activo, TipoDispositivo
from apps.activos.services import ActivoService
from apps.mantenimientos.models import CatalogoComponente, Mantenimiento
from apps.mantenimientos.services import MantenimientoService
from apps.organizacion.models import Departamento, Empleado
from apps.politicas.models import PoliticaObsolescencia
from apps.users.models import User


@pytest.fixture
def admin(db):
    return User.objects.create_superuser(
        username="admin_dash", email="admin_dash@example.com", password="Sup3r-Secr3t!"
    )


@pytest.fixture
def cliente(admin):
    client = APIClient()
    client.force_authenticate(user=admin)
    return client


@pytest.fixture
def escenario(db, admin):
    """Inventario pequeño pero con un activo en cada estado relevante."""
    departamento = Departamento.objects.create(nombre="Tecnología", codigo="TI")
    tipo = TipoDispositivo.objects.create(nombre="Laptop", codigo="LAP")
    empleado = Empleado.objects.create(
        nombres="Ana", apellidos="Pérez", codigo_empleado="EMP-0001", departamento=departamento
    )

    def crear(serie, fecha=datetime.date(2024, 1, 10)):
        return ActivoService.crear_activo(
            actor=admin,
            tipo=tipo,
            nombre=f"Equipo {serie}",
            marca="Dell",
            modelo="Latitude",
            numero_serie=serie,
            departamento=departamento,
            fecha_adquisicion=fecha,
            costo_adquisicion=Decimal("1000.00"),
        )

    en_uso = crear("SN-DASH-1")
    ActivoService.asignar_responsables(actor=admin, activo=en_uso, responsables=[empleado])
    en_bodega = crear("SN-DASH-2")
    de_baja = crear("SN-DASH-3")
    ActivoService.cambiar_estado(
        actor=admin, activo=de_baja, estado=Activo.Estado.DADO_DE_BAJA, motivo="Obsoleto"
    )

    return {
        "departamento": departamento,
        "tipo": tipo,
        "empleado": empleado,
        "en_uso": en_uso,
        "en_bodega": en_bodega,
        "de_baja": de_baja,
        "crear": crear,
    }


def _registrar_mantenimiento(admin, activo, fecha, costo=Decimal("25.00"), componentes=None):
    return MantenimientoService.registrar(
        actor=admin,
        activo=activo,
        componentes=componentes,
        tipo=Mantenimiento.Tipo.CORRECTIVO,
        fecha_intervencion=fecha,
        tipo_responsable=Mantenimiento.TipoResponsable.TECNICO_INTERNO,
        responsable="Soporte",
        descripcion="Reparación",
        costo_mano_obra=costo,
    )


# --- Dashboard -------------------------------------------------------------


def test_el_dashboard_cuenta_los_activos_por_estado(cliente, escenario):
    datos = cliente.get("/api/v1/activos/dashboard/").json()

    assert datos["activos"]["total"] == 3
    assert datos["activos"]["en_uso"] == 1
    assert datos["activos"]["en_bodega"] == 1
    assert datos["activos"]["dados_de_baja"] == 1


def test_el_dashboard_informa_el_costo_acumulado_de_reparaciones(cliente, admin, escenario):
    from django.utils import timezone

    disco = CatalogoComponente.objects.create(nombre="Disco", codigo="SSD", es_critico=True)
    _registrar_mantenimiento(
        admin,
        escenario["en_uso"],
        timezone.localdate(),
        costo=Decimal("30.00"),
        componentes=[{"componente": disco, "cantidad": 2, "costo_unitario": Decimal("50.00")}],
    )

    datos = cliente.get("/api/v1/activos/dashboard/").json()

    # 30 de mano de obra + 2 x 50 de repuestos.
    assert Decimal(datos["mantenimientos"]["costo_acumulado"]) == Decimal("130.00")
    assert datos["mantenimientos"]["del_mes"] == 1


def test_las_reparaciones_del_mes_no_cuentan_las_de_meses_anteriores(cliente, admin, escenario):
    from django.utils import timezone

    hoy = timezone.localdate()
    _registrar_mantenimiento(admin, escenario["en_uso"], hoy)
    _registrar_mantenimiento(admin, escenario["en_bodega"], datetime.date(2024, 6, 1))

    datos = cliente.get("/api/v1/activos/dashboard/").json()

    assert datos["mantenimientos"]["del_mes"] == 1
    assert datos["mantenimientos"]["total_historico"] == 2


def test_el_ranking_ordena_por_cantidad_de_intervenciones(cliente, admin, escenario):
    for _ in range(3):
        _registrar_mantenimiento(admin, escenario["en_uso"], datetime.date(2024, 6, 1))
    _registrar_mantenimiento(admin, escenario["en_bodega"], datetime.date(2024, 6, 1))

    ranking = cliente.get("/api/v1/activos/dashboard/").json()["equipos_mas_reparados"]

    assert ranking[0]["codigo_barras"] == escenario["en_uso"].codigo_barras
    assert ranking[0]["total_mantenimientos"] == 3


def test_el_ranking_omite_los_activos_dados_de_baja(cliente, admin, escenario):
    """Un equipo retirado no es un problema a resolver: sesgaría el ranking de
    los que sí siguen operando."""
    _registrar_mantenimiento(admin, escenario["de_baja"], datetime.date(2024, 6, 1))

    ranking = cliente.get("/api/v1/activos/dashboard/").json()["equipos_mas_reparados"]

    assert all(e["codigo_barras"] != escenario["de_baja"].codigo_barras for e in ranking)


def test_el_dashboard_cuenta_los_activos_con_sugerencia_de_renovacion(cliente, admin, escenario):
    PoliticaObsolescencia.objects.create(
        nombre="Laptops", tipo_dispositivo=escenario["tipo"], max_mantenimientos=0
    )
    _registrar_mantenimiento(admin, escenario["en_uso"], datetime.date(2024, 6, 1))

    datos = cliente.get("/api/v1/activos/dashboard/").json()

    assert datos["activos"]["requieren_renovacion"] == 1


def test_el_dashboard_cubre_los_diez_indicadores_del_documento(cliente, escenario):
    """Ya no queda ninguno sin calcular; la clave se conserva para poder
    declarar los que aparezcan en el futuro."""
    datos = cliente.get("/api/v1/activos/dashboard/").json()

    assert datos["indicadores_no_disponibles"] == []
    assert "garantias" in datos


def test_el_dashboard_distingue_garantias_vencidas_de_no_registradas(cliente, escenario):
    """No es lo mismo una cobertura expirada que una fecha nunca capturada:
    mezclarlas haría que un inventario a medio llenar pareciera un parque
    entero fuera de garantía."""
    from datetime import timedelta

    from django.utils import timezone

    hoy = timezone.localdate()
    vencido = escenario["crear"]("SN-GAR-1")
    vencido.fecha_fin_garantia = hoy - timedelta(days=1)
    vencido.save()

    por_vencer = escenario["crear"]("SN-GAR-2")
    por_vencer.fecha_fin_garantia = hoy + timedelta(days=10)
    por_vencer.save()

    vigente = escenario["crear"]("SN-GAR-3")
    vigente.fecha_fin_garantia = hoy + timedelta(days=365)
    vigente.save()

    garantias = cliente.get("/api/v1/activos/dashboard/").json()["garantias"]

    assert garantias["vencidas"] == 1
    assert garantias["por_vencer"] == 1
    assert garantias["vigentes"] == 1
    # Dos del escenario base: el tercero está dado de baja, y el panel muestra
    # el parque operativo (un equipo retirado no tiene garantía que gestionar).
    assert garantias["sin_registrar"] == 2


def test_el_dashboard_informa_el_tiempo_fuera_de_operacion(cliente, admin, escenario):
    """Solo cuenta las intervenciones cerradas: mientras no haya fecha de
    salida, el equipo sigue fuera y ese tiempo aún no está determinado."""
    cerrada = _registrar_mantenimiento(admin, escenario["en_uso"], datetime.date(2024, 6, 1))
    cerrada.fecha_salida = datetime.date(2024, 6, 4)
    cerrada.save()
    _registrar_mantenimiento(admin, escenario["en_bodega"], datetime.date(2024, 7, 1))

    fuera = cliente.get("/api/v1/activos/dashboard/").json()["fuera_de_operacion"]

    assert fuera["total_dias"] == 3
    assert fuera["intervenciones_cerradas"] == 1
    assert fuera["intervenciones_abiertas"] == 1


def test_el_dashboard_agrupa_por_tipo_y_por_departamento(cliente, escenario):
    datos = cliente.get("/api/v1/activos/dashboard/").json()

    por_tipo = {fila["nombre_tipo"]: fila["total"] for fila in datos["por_tipo_dispositivo"]}
    por_area = {fila["nombre_departamento"]: fila for fila in datos["por_departamento"]}

    # Los dados de baja quedan fuera: el panel muestra el parque operativo.
    assert por_tipo["Laptop"] == 2
    assert por_area["Tecnología"]["total"] == 2
    assert por_area["Tecnología"]["asignados"] == 1


# --- Exportación -----------------------------------------------------------


def test_el_inventario_se_exporta_a_excel(cliente, escenario):
    respuesta = cliente.get("/api/v1/activos/exportar/")

    assert respuesta.status_code == 200
    assert "spreadsheetml" in respuesta["Content-Type"]
    assert "inventario-activos.xlsx" in respuesta["Content-Disposition"]

    hoja = load_workbook(BytesIO(respuesta.content))["Inventario"]
    encabezados = [c.value for c in hoja[1]]
    assert "Código de barras" in encabezados
    assert "Sugerencia de renovación" in encabezados
    assert hoja.max_row == 4  # cabecera + 3 activos


def test_la_exportacion_respeta_los_filtros_del_listado(cliente, escenario):
    """El archivo trae lo que el usuario está viendo: exportar siempre el
    inventario completo le obligaría a filtrar otra vez en Excel."""
    respuesta = cliente.get("/api/v1/activos/exportar/?estado=en_uso")

    hoja = load_workbook(BytesIO(respuesta.content))["Inventario"]

    assert hoja.max_row == 2  # cabecera + 1 activo
    assert hoja.cell(2, 1).value == escenario["en_uso"].codigo_barras


def test_la_exportacion_incluye_custodio_y_especificaciones(cliente, admin, escenario):
    escenario["en_uso"].especificaciones = {"RAM": "16 GB"}
    escenario["en_uso"].save()

    respuesta = cliente.get("/api/v1/activos/exportar/?estado=en_uso")
    hoja = load_workbook(BytesIO(respuesta.content))["Inventario"]
    fila = [c.value for c in hoja[2]]

    assert "Ana Pérez" in fila
    assert "RAM=16 GB" in fila


def test_la_bitacora_de_mantenimientos_se_exporta_a_excel(cliente, admin, escenario):
    disco = CatalogoComponente.objects.create(nombre="Disco", codigo="SSD", es_critico=True)
    _registrar_mantenimiento(
        admin,
        escenario["en_uso"],
        datetime.date(2024, 6, 1),
        costo=Decimal("30.00"),
        componentes=[{"componente": disco, "cantidad": 1, "costo_unitario": Decimal("50.00")}],
    )

    respuesta = cliente.get("/api/v1/mantenimientos/exportar/")

    assert respuesta.status_code == 200
    hoja = load_workbook(BytesIO(respuesta.content))["Mantenimientos"]
    fila = [c.value for c in hoja[2]]
    assert escenario["en_uso"].codigo_barras in fila
    assert "Disco x1 (crítica)" in fila
    assert Decimal(str(fila[-1])) == Decimal("80.00")


def test_exportar_exige_su_propio_permiso(db, escenario):
    """Mismo criterio que `auditoria.exportar`: sacar el inventario completo en
    un archivo es distinto de consultarlo en pantalla."""
    from django.contrib.auth.models import Group, Permission
    from django.contrib.contenttypes.models import ContentType

    from apps.permissions.models import ModulePermission

    usuario = User.objects.create_user(
        username="solo_consulta", email="sc@example.com", password="Sup3r-Secr3t!"
    )
    grupo = Group.objects.create(name="Solo consulta")
    grupo.permissions.set(
        Permission.objects.filter(
            content_type=ContentType.objects.get_for_model(ModulePermission),
            codename__in=["activos.ver", "mantenimientos.ver"],
        )
    )
    usuario.groups.add(grupo)
    client = APIClient()
    client.force_authenticate(user=usuario)

    assert client.get("/api/v1/activos/").status_code == 200
    assert client.get("/api/v1/activos/dashboard/").status_code == 200
    assert client.get("/api/v1/activos/exportar/").status_code == 403
    assert client.get("/api/v1/mantenimientos/exportar/").status_code == 403


def test_la_exportacion_queda_auditada(cliente, escenario):
    from apps.core.models import AuditLog

    cliente.get("/api/v1/activos/exportar/?estado=en_uso")

    evento = AuditLog.objects.filter(action="activo.exportado").get()
    assert evento.new_values["filas"] == 1
    assert evento.new_values["filtros"]["estado"] == ["en_uso"]
