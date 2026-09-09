"""Carga masiva de activos desde .xlsx: validación previa e importación."""

import datetime
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from apps.activos.importacion import MAX_FILAS, columnas_configuradas
from apps.activos.models import Activo, TipoDispositivo
from apps.activos.services import ActivoService
from apps.organizacion.models import Departamento, Empleado
from apps.users.models import User


@pytest.fixture
def columnas(db):
    """Columnas vigentes de la plantilla, sembradas por migración.

    Se leen en cada prueba en vez de fijarlas como constante del módulo: la
    plantilla es configurable, y una lista fija haría que las pruebas dejaran
    de verificar lo que el sistema hace de verdad en cuanto alguien la ajuste.
    """
    return columnas_configuradas()


def _encabezados(columnas):
    return [columna.encabezado for columna in columnas]


def _claves(columnas):
    return [columna.clave for columna in columnas]


@pytest.fixture
def admin(db):
    return User.objects.create_superuser(
        username="admin_import", email="admin_import@example.com", password="Sup3r-Secr3t!"
    )


@pytest.fixture
def cliente(admin):
    client = APIClient()
    client.force_authenticate(user=admin)
    return client


@pytest.fixture
def catalogos(db):
    departamento = Departamento.objects.create(nombre="Tecnología", codigo="TI")
    tipo = TipoDispositivo.objects.create(nombre="Laptop", codigo="LAP")
    empleado = Empleado.objects.create(
        nombres="Ana",
        apellidos="Pérez",
        codigo_empleado="EMP-0001",
        departamento=departamento,
    )
    return {"departamento": departamento, "tipo": tipo, "empleado": empleado}


def _fila(**overrides):
    base = {
        "tipo": "LAP",
        "nombre": "Laptop 01",
        "marca": "Dell",
        "modelo": "Latitude",
        "numero_serie": "SN-IMP-1",
        "departamento": "TI",
        "fecha_adquisicion": "2024-01-15",
        "custodio": "",
        "sede": "",
        "costo_adquisicion": "",
        "especificaciones": "",
        "observaciones": "",
    }
    base.update(overrides)
    return base


def _archivo(columnas, filas, encabezados=None):
    """Construye un .xlsx en memoria con los encabezados de la plantilla."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Activos"
    hoja.append(encabezados if encabezados is not None else _encabezados(columnas))
    for fila in filas:
        hoja.append([fila.get(clave, "") for clave in _claves(columnas)])
    buffer = BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    buffer.name = "carga.xlsx"
    return buffer


def _importar(cliente, archivo, confirmar=False):
    return cliente.post(
        "/api/v1/activos/importar/",
        {"archivo": archivo, "confirmar": "true" if confirmar else "false"},
        format="multipart",
    )


# --- Plantilla -------------------------------------------------------------


def test_la_plantilla_se_descarga_como_xlsx(cliente, catalogos):
    respuesta = cliente.get("/api/v1/activos/plantilla-importacion/")

    assert respuesta.status_code == 200
    assert "spreadsheetml" in respuesta["Content-Type"]
    assert "plantilla-carga-activos.xlsx" in respuesta["Content-Disposition"]


def test_la_plantilla_trae_solo_lo_de_activos(cliente, catalogos):
    """Llegó a tener ocho hojas, cinco de ellas solo de consulta: mezclaba
    material de referencia con material de trabajo, y ninguno de esos catálogos
    se podía cargar. Cada uno tiene ahora su propio archivo."""
    respuesta = cliente.get("/api/v1/activos/plantilla-importacion/")
    libro = load_workbook(BytesIO(respuesta.content))

    assert libro.sheetnames == ["Instrucciones", "Activos", "Ejemplo"]


def test_las_instrucciones_dicen_donde_estan_los_valores_validos(cliente, catalogos):
    """Quitar las hojas de catálogo sin decir a dónde fueron dejaría al usuario
    adivinando qué escribir en «Tipo de dispositivo»."""
    respuesta = cliente.get("/api/v1/activos/plantilla-importacion/")
    hoja = load_workbook(BytesIO(respuesta.content))["Instrucciones"]
    texto = " ".join(str(fila[0].value or "") for fila in hoja.iter_rows())

    assert "Carga masiva" in texto
    assert "plantillas" in texto


def test_la_hoja_de_captura_queda_vacia_bajo_los_encabezados(cliente, catalogos, columnas):
    """Sin filas de ejemplo que haya que acordarse de borrar: olvidarlas
    produciría filas basura en el inventario."""
    respuesta = cliente.get("/api/v1/activos/plantilla-importacion/")
    hoja = load_workbook(BytesIO(respuesta.content))["Activos"]

    assert [c.value for c in hoja[1]] == _encabezados(columnas)
    assert hoja.max_row == 1


def test_la_plantilla_se_puede_llenar_y_cargar_sin_ajustes(cliente, catalogos, columnas):
    """La plantilla descargada debe ser directamente utilizable: si sus
    encabezados no coincidieran con los que espera el importador, el usuario
    recibiría un error incomprensible."""
    respuesta = cliente.get("/api/v1/activos/plantilla-importacion/")
    libro = load_workbook(BytesIO(respuesta.content))
    hoja = libro["Activos"]
    hoja.append([_fila()[clave] for clave in _claves(columnas)])

    buffer = BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    buffer.name = "carga.xlsx"

    reporte = _importar(cliente, buffer).json()

    assert reporte["errores"] == []
    assert reporte["filas_validas"] == 1


# --- Validación previa -----------------------------------------------------


def test_la_validacion_no_escribe_nada_en_el_inventario(cliente, catalogos, columnas):
    _importar(cliente, _archivo(columnas, [_fila(), _fila(numero_serie="SN-IMP-2")]))

    assert Activo.objects.count() == 0


def test_el_reporte_indica_cuantas_filas_son_validas(cliente, catalogos, columnas):
    respuesta = _importar(cliente, _archivo(columnas, [_fila(), _fila(numero_serie="SN-IMP-2")]))

    reporte = respuesta.json()
    assert reporte["total_filas"] == 2
    assert reporte["filas_validas"] == 2
    assert reporte["es_importable"] is True
    assert reporte["importado"] is False


def test_una_serie_repetida_dentro_del_archivo_se_detecta(cliente, catalogos, columnas):
    respuesta = _importar(cliente, _archivo(columnas, [_fila(), _fila(nombre="Otra")]))

    errores = respuesta.json()["errores"]
    assert any("repetida" in e["mensaje"] for e in errores)
    assert respuesta.json()["es_importable"] is False


def test_una_serie_ya_existente_en_el_inventario_se_detecta(cliente, admin, catalogos, columnas):
    ActivoService.crear_activo(
        actor=admin,
        tipo=catalogos["tipo"],
        nombre="Existente",
        marca="Dell",
        modelo="X",
        numero_serie="SN-IMP-1",
        departamento=catalogos["departamento"],
        fecha_adquisicion=datetime.date(2024, 1, 1),
    )

    respuesta = _importar(cliente, _archivo(columnas, [_fila()]))

    assert any("Ya existe" in e["mensaje"] for e in respuesta.json()["errores"])


def test_un_tipo_o_area_inexistente_se_reporta_con_la_fila_y_la_columna(
    cliente, catalogos, columnas
):
    respuesta = _importar(
        cliente,
        _archivo(columnas, [_fila(tipo="NOEXISTE"), _fila(numero_serie="SN-2", departamento="XX")]),
    )

    errores = respuesta.json()["errores"]
    tipo_error = next(e for e in errores if e["columna"] == "Tipo de dispositivo")
    area_error = next(e for e in errores if e["columna"] == "Departamento")
    assert tipo_error["fila"] == 2  # fila 1 = encabezados
    assert area_error["fila"] == 3


def test_una_fecha_mal_escrita_se_reporta(cliente, catalogos, columnas):
    respuesta = _importar(cliente, _archivo(columnas, [_fila(fecha_adquisicion="15/13/2024")]))

    assert any(e["columna"] == "Fecha de adquisición" for e in respuesta.json()["errores"])


def test_una_fecha_futura_se_rechaza(cliente, catalogos, columnas):
    manana = datetime.date.today() + datetime.timedelta(days=1)

    respuesta = _importar(
        cliente, _archivo(columnas, [_fila(fecha_adquisicion=manana.isoformat())])
    )

    assert any("futura" in e["mensaje"] for e in respuesta.json()["errores"])


def test_un_custodio_inexistente_se_reporta(cliente, catalogos, columnas):
    respuesta = _importar(cliente, _archivo(columnas, [_fila(custodio="NO-EXISTE")]))

    assert any(e["columna"] == "Código del custodio" for e in respuesta.json()["errores"])


def test_un_costo_con_simbolo_de_moneda_se_reporta(cliente, catalogos, columnas):
    respuesta = _importar(cliente, _archivo(columnas, [_fila(costo_adquisicion="$1.150,00")]))

    assert any(e["columna"] == "Costo de compra" for e in respuesta.json()["errores"])


def test_faltar_una_columna_obligatoria_lo_dice_en_vez_de_fallar_fila_por_fila(
    cliente, catalogos, columnas
):
    encabezados_incompletos = [e for e in _encabezados(columnas) if e != "Número de serie *"]

    respuesta = _importar(cliente, _archivo(columnas, [], encabezados=encabezados_incompletos))

    error = respuesta.json()["errores"][0]
    assert error["columna"] == "encabezados"
    assert "Número de serie" in error["mensaje"]


def test_las_filas_vacias_del_final_no_cuentan_como_datos(cliente, catalogos, columnas):
    """Excel deja filas en blanco al final de una hoja editada; tomarlas como
    activos produciría un error por cada una."""
    archivo = _archivo(columnas, [_fila()])
    libro = load_workbook(archivo)
    hoja = libro.active
    for _ in range(5):
        hoja.append(["" for _ in _claves(columnas)])
    buffer = BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    buffer.name = "carga.xlsx"

    reporte = _importar(cliente, buffer).json()

    assert reporte["total_filas"] == 1
    assert reporte["errores"] == []


# --- Importación -----------------------------------------------------------


def test_confirmar_crea_los_activos_con_su_codigo_de_barras(cliente, catalogos, columnas):
    respuesta = _importar(
        cliente,
        _archivo(
            columnas,
            [
                _fila(custodio="EMP-0001", costo_adquisicion="1150.50"),
                _fila(numero_serie="SN-IMP-2", nombre="Laptop 02"),
            ],
        ),
        confirmar=True,
    )

    assert respuesta.status_code == 201
    assert respuesta.json()["importado"] is True
    assert Activo.objects.count() == 2
    codigos = sorted(Activo.objects.values_list("codigo_barras", flat=True))
    assert codigos == ["GA-LAP-000001", "GA-LAP-000002"]


def test_la_importacion_deja_el_movimiento_de_alta_de_cada_activo(cliente, catalogos, columnas):
    _importar(cliente, _archivo(columnas, [_fila()]), confirmar=True)

    activo = Activo.objects.get()
    assert activo.movimientos.count() == 1
    assert activo.movimientos.get().tipo == "alta"


def test_un_archivo_con_errores_no_importa_ninguna_fila(cliente, catalogos, columnas):
    """Todo o nada: un inventario a medio cargar es peor que uno vacío, porque
    nadie sabe cuál de los dos casos está mirando."""
    respuesta = _importar(
        cliente,
        _archivo(columnas, [_fila(), _fila(numero_serie="SN-IMP-2", tipo="NOEXISTE")]),
        confirmar=True,
    )

    assert respuesta.status_code == 400
    assert respuesta.json()["importado"] is False
    assert Activo.objects.count() == 0


def test_las_especificaciones_se_convierten_en_pares_clave_valor(cliente, catalogos, columnas):
    _importar(
        cliente,
        _archivo(columnas, [_fila(especificaciones="Procesador=i5; RAM=16 GB; Disco=512 GB SSD")]),
        confirmar=True,
    )

    activo = Activo.objects.get()
    assert activo.especificaciones == {
        "Procesador": "i5",
        "RAM": "16 GB",
        "Disco": "512 GB SSD",
    }


def test_un_activo_sin_custodio_queda_en_bodega(cliente, catalogos, columnas):
    _importar(cliente, _archivo(columnas, [_fila(custodio="")]), confirmar=True)

    activo = Activo.objects.get()
    assert activo.custodio is None
    assert activo.estado == Activo.Estado.EN_BODEGA


def test_la_importacion_queda_auditada(cliente, catalogos, columnas):
    from apps.core.models import AuditLog

    _importar(cliente, _archivo(columnas, [_fila()]), confirmar=True)

    evento = AuditLog.objects.filter(action="activo.importacion_masiva").get()
    assert evento.new_values["cantidad"] == 1
    assert evento.new_values["codigos_generados"] == ["GA-LAP-000001"]


# --- Validaciones del archivo en sí ---------------------------------------


def test_se_rechaza_un_archivo_que_no_sea_xlsx(cliente, catalogos):
    from django.core.files.uploadedfile import SimpleUploadedFile

    archivo = SimpleUploadedFile("datos.csv", b"tipo,nombre\nLAP,x", content_type="text/csv")

    respuesta = cliente.post("/api/v1/activos/importar/", {"archivo": archivo}, format="multipart")

    assert respuesta.status_code == 400
    assert respuesta.json()["error"]["code"] == "formato_invalido"


def test_un_xlsx_corrupto_da_un_mensaje_util_y_no_una_traza(cliente, catalogos):
    from django.core.files.uploadedfile import SimpleUploadedFile

    archivo = SimpleUploadedFile(
        "roto.xlsx", b"esto no es un excel", content_type="application/vnd.ms-excel"
    )

    respuesta = cliente.post("/api/v1/activos/importar/", {"archivo": archivo}, format="multipart")

    assert respuesta.status_code == 400
    assert respuesta.json()["error"]["code"] == "archivo_ilegible"


def test_sin_archivo_adjunto_se_pide_uno(cliente, catalogos):
    respuesta = cliente.post("/api/v1/activos/importar/", {}, format="multipart")

    assert respuesta.status_code == 400
    assert respuesta.json()["error"]["code"] == "archivo_requerido"


def test_se_rechaza_un_archivo_con_mas_filas_del_maximo(cliente, catalogos, columnas):
    filas = [_fila(numero_serie=f"SN-MASIVO-{n}") for n in range(MAX_FILAS + 5)]

    respuesta = _importar(cliente, _archivo(columnas, filas))

    assert any("máximo" in e["mensaje"] for e in respuesta.json()["errores"])


def test_importar_exige_el_permiso_de_crear_activos(db, catalogos, columnas):
    """Ver el inventario no alcanza para cargarlo masivamente."""
    from django.contrib.auth.models import Group, Permission
    from django.contrib.contenttypes.models import ContentType

    from apps.permissions.models import ModulePermission

    usuario = User.objects.create_user(
        username="solo_lectura_import", email="sl@example.com", password="Sup3r-Secr3t!"
    )
    grupo = Group.objects.create(name="Solo lectura import")
    grupo.permissions.set(
        Permission.objects.filter(
            content_type=ContentType.objects.get_for_model(ModulePermission),
            codename="activos.ver",
        )
    )
    usuario.groups.add(grupo)

    client = APIClient()
    client.force_authenticate(user=usuario)

    assert client.get("/api/v1/activos/plantilla-importacion/").status_code == 403
    assert _importar(client, _archivo(columnas, [_fila()])).status_code == 403


# --- Nombres repetidos: se avisan, no bloquean ------------------------------


def test_el_nombre_repetido_dentro_del_archivo_avisa_sin_bloquear(cliente, catalogos, columnas):
    """El nombre no identifica: dos equipos pueden llamarse «Laptop 01» sin que
    nada esté mal. Repetirlo sí deja un inventario que no se lee de un vistazo
    —dos filas idénticas salvo la serie—, y sobre todo delata la fila pegada
    dos veces, que es el accidente habitual al armar la plantilla."""
    archivo = _archivo(
        columnas,
        [_fila(numero_serie="SN-DUP-1"), _fila(numero_serie="SN-DUP-2")],
    )

    reporte = _importar(cliente, archivo).json()

    assert reporte["errores"] == []
    assert reporte["es_importable"] is True
    assert len(reporte["advertencias"]) == 1
    aviso = reporte["advertencias"][0]
    assert aviso["fila"] == 3
    assert "fila 2" in aviso["mensaje"]


def test_el_nombre_que_ya_existe_en_el_inventario_tambien_avisa(
    cliente, catalogos, columnas, admin
):
    ActivoService.crear_activo(
        actor=admin,
        tipo=catalogos["tipo"],
        nombre="Laptop 01",
        marca="Dell",
        modelo="Latitude",
        numero_serie="SN-YA-EXISTE",
        departamento=catalogos["departamento"],
        fecha_adquisicion=datetime.date(2024, 1, 15),
    )

    reporte = _importar(cliente, _archivo(columnas, [_fila(numero_serie="SN-NUEVA")])).json()

    assert reporte["es_importable"] is True
    assert len(reporte["advertencias"]) == 1
    assert "Ya hay un activo llamado" in reporte["advertencias"][0]["mensaje"]


def test_la_serie_repetida_si_bloquea(cliente, catalogos, columnas):
    """La diferencia con el nombre: la serie identifica al equipo físico, y dos
    activos con la misma serie hacen que escanearla devuelva el equivocado."""
    archivo = _archivo(columnas, [_fila(), _fila(nombre="Laptop 02")])

    reporte = _importar(cliente, archivo).json()

    assert reporte["es_importable"] is False
    assert reporte["filas_con_error"] == 1
    assert "repetida en la fila 2" in reporte["errores"][0]["mensaje"]


def test_no_se_avisa_de_una_fila_que_ademas_no_entra(cliente, catalogos, columnas):
    """Las advertencias de una fila que no va a importarse sobran: lo que hay
    que mirar es por qué no entra."""
    archivo = _archivo(
        columnas,
        [
            _fila(numero_serie="SN-OK"),
            _fila(numero_serie="SN-MALA", tipo="NO-EXISTE"),
        ],
    )

    reporte = _importar(cliente, archivo).json()

    assert reporte["errores"] != []
    assert reporte["advertencias"] == []


def test_un_mismo_problema_no_se_reporta_dos_veces(cliente, catalogos, columnas):
    """Un campo obligatorio que además no está en el catálogo lo señalan dos
    comprobaciones distintas. Repetido en el reporte parece que hay dos
    problemas donde hay uno."""
    archivo = _archivo(columnas, [_fila(departamento="")])

    reporte = _importar(cliente, archivo).json()

    del_departamento = [e for e in reporte["errores"] if e["columna"] == "Departamento"]
    assert len(del_departamento) == 1


def test_el_aviso_senala_la_fila_del_archivo_y_no_el_activo_de_la_base(
    cliente, catalogos, columnas, admin
):
    """Si el nombre ya existía en el inventario, la segunda fila del archivo
    volvía a señalar al activo de la base y no a la fila de al lado, que es la
    que hay que mirar."""
    ActivoService.crear_activo(
        actor=admin,
        tipo=catalogos["tipo"],
        nombre="Laptop 01",
        marca="Dell",
        modelo="Latitude",
        numero_serie="SN-BASE",
        departamento=catalogos["departamento"],
        fecha_adquisicion=datetime.date(2024, 1, 15),
    )
    archivo = _archivo(
        columnas,
        [_fila(numero_serie="SN-A"), _fila(numero_serie="SN-B")],
    )

    avisos = _importar(cliente, archivo).json()["advertencias"]

    assert "Ya hay un activo llamado" in avisos[0]["mensaje"]
    assert "fila 2 de este mismo archivo" in avisos[1]["mensaje"]
