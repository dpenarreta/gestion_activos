"""Exportación del inventario a Excel (§16 del documento funcional).

Exporta **lo que el usuario está viendo**: recibe el queryset ya filtrado por
la vista, así que el archivo refleja los mismos filtros que la pantalla. Un
export que siempre baja el inventario completo obliga a filtrar otra vez en
Excel y hace inútil el trabajo que el usuario ya hizo en la interfaz.

Se escribe en modo `write_only`, que vuelca fila por fila sin mantener la hoja
entera en memoria: con las 5.000 a 10.000 filas que dimensiona el documento, la
diferencia es notable.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.core.hojas_de_calculo import neutralizar_fila

from .models import Activo

RELLENO_CABECERA = PatternFill("solid", fgColor="D9E2F3")

#: Se exporta el nivel y no un sí/no: «Reemplazo recomendado» y «Evaluar
#: reemplazo» se presupuestan distinto, y en la hoja de cálculo ese matiz
#: es justo lo que se filtra.
NIVELES_RENOVACION = {
    "ninguno": "No",
    "evaluar": "Evaluar reemplazo",
    "recomendado": "Reemplazo recomendado",
}

COLUMNAS_ACTIVOS = [
    ("Código de barras", 18),
    ("Nombre", 30),
    ("Tipo", 16),
    ("Marca", 16),
    ("Modelo", 20),
    ("Número de serie", 22),
    ("Estado", 18),
    ("Responsables", 30),
    ("Departamento", 20),
    ("Sede", 20),
    ("Ciudad", 16),
    ("Criticidad", 12),
    ("Uso", 18),
    ("Fecha de adquisición", 18),
    ("Fecha de ingreso", 18),
    ("Antigüedad (meses)", 16),
    ("Costo de compra", 16),
    ("Condición", 12),
    ("Proveedor", 24),
    ("Propiedad", 16),
    ("Concesionario", 24),
    ("Fin de garantía", 16),
    ("Estado de garantía", 20),
    ("Mantenimientos", 14),
    ("Piezas críticas", 14),
    ("Sugerencia de renovación", 24),
    ("Especificaciones", 40),
    ("Observaciones", 30),
]

COLUMNAS_MANTENIMIENTOS = [
    ("Ingreso", 12),
    ("Salida", 12),
    ("Días fuera", 12),
    ("Código del activo", 18),
    ("Activo", 28),
    ("Tipo", 14),
    ("Responsable", 24),
    ("A cargo de", 18),
    ("Causa", 24),
    ("Trabajo realizado", 40),
    ("Diagnóstico", 30),
    ("Solución", 30),
    ("Estado final", 16),
    ("Garantía usada", 14),
    ("Componentes", 34),
    ("Mano de obra", 14),
    ("Repuestos", 14),
    ("Costo total", 14),
]


def _hoja_con_formato(libro, titulo, columnas):
    from openpyxl.cell import WriteOnlyCell

    hoja = libro.create_sheet(titulo)
    cabecera = []
    for indice, (nombre, ancho) in enumerate(columnas, start=1):
        celda = WriteOnlyCell(hoja, value=nombre)
        celda.font = Font(bold=True)
        celda.fill = RELLENO_CABECERA
        celda.alignment = Alignment(vertical="center", wrap_text=True)
        cabecera.append(celda)
        hoja.column_dimensions[get_column_letter(indice)].width = ancho
    hoja.append(cabecera)
    hoja.freeze_panes = "A2"
    return hoja


def _especificaciones_a_texto(especificaciones: dict) -> str:
    return "; ".join(f"{clave}={valor}" for clave, valor in (especificaciones or {}).items())


def exportar_activos(queryset) -> bytes:
    """Inventario en .xlsx, respetando el filtro del queryset recibido."""
    libro = Workbook(write_only=True)
    hoja = _hoja_con_formato(libro, "Inventario", COLUMNAS_ACTIVOS)

    consulta = queryset.select_related(
        "tipo", "departamento", "sede", "proveedor", "concesionario"
    ).prefetch_related("responsables")
    for activo in consulta.iterator(chunk_size=500):
        hoja.append(
            neutralizar_fila(
                [
                    activo.codigo_barras,
                    activo.nombre,
                    activo.tipo.nombre,
                    activo.marca,
                    activo.modelo,
                    activo.numero_serie,
                    activo.get_estado_display(),
                    activo.resumen_de_responsables,
                    activo.departamento.nombre,
                    activo.sede.nombre if activo.sede_id else "",
                    activo.sede.donde if activo.sede_id else "",
                    activo.get_criticidad_display(),
                    activo.get_uso_display(),
                    activo.fecha_adquisicion,
                    activo.fecha_ingreso,
                    activo.antiguedad_meses,
                    activo.costo_adquisicion,
                    activo.get_condicion_display(),
                    activo.proveedor.nombre if activo.proveedor_id else "",
                    activo.get_propiedad_display(),
                    activo.concesionario.nombre if activo.concesionario_id else "",
                    activo.fecha_fin_garantia,
                    Activo.Garantia(activo.estado_garantia).label,
                    activo.total_mantenimientos,
                    activo.total_componentes_criticos,
                    NIVELES_RENOVACION.get(activo.nivel_renovacion, "No"),
                    _especificaciones_a_texto(activo.especificaciones),
                    activo.observaciones,
                ]
            )
        )

    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def exportar_mantenimientos(queryset) -> bytes:
    """Bitácora de mantenimientos en .xlsx, con el desglose de repuestos."""
    from decimal import Decimal

    libro = Workbook(write_only=True)
    hoja = _hoja_con_formato(libro, "Mantenimientos", COLUMNAS_MANTENIMIENTOS)

    consulta = queryset.select_related("activo").prefetch_related("componentes__componente")
    for mantenimiento in consulta.iterator(chunk_size=500):
        componentes = list(mantenimiento.componentes.all())
        repuestos = sum(
            (c.costo_unitario or Decimal("0")) * c.cantidad for c in componentes
        ) or Decimal("0")
        mano_obra = mantenimiento.costo_mano_obra or Decimal("0")

        hoja.append(
            neutralizar_fila(
                [
                    mantenimiento.fecha_intervencion,
                    mantenimiento.fecha_salida,
                    mantenimiento.dias_fuera_de_operacion,
                    mantenimiento.activo.codigo_barras,
                    mantenimiento.activo.nombre,
                    mantenimiento.get_tipo_display(),
                    mantenimiento.responsable,
                    mantenimiento.get_tipo_responsable_display(),
                    mantenimiento.causa,
                    mantenimiento.descripcion,
                    mantenimiento.diagnostico,
                    mantenimiento.solucion,
                    mantenimiento.get_estado_final_display(),
                    "Sí" if mantenimiento.garantia_usada else "No",
                    "; ".join(
                        f"{c.componente.nombre} x{c.cantidad}"
                        + (" (crítica)" if c.era_critico else "")
                        for c in componentes
                    ),
                    mano_obra,
                    repuestos,
                    mano_obra + repuestos,
                ]
            )
        )

    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()
