"""Matriz de obligaciones, fase 5: generación de períodos, calendario,
reportes, exportaciones, auditoría y membresías.

Cobertura de tests/qa/features/matriz-obligaciones.feature (AC-MTZ-047 a
AC-MTZ-056). Sigue las pruebas mínimas de la especificación (§23):
generación mensual, trimestral y anual, idempotencia, única no genera;
reportes con porcentajes, por área y entidad y alcance propio.
"""

import datetime
import io
import zoneinfo

import pytest
from django.contrib.auth.models import Group
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.core.models import AuditLog
from apps.empresas.models import Empresa, MembresiaEmpresa
from apps.empresas.servicios import asignar_empresas
from apps.matriz.generacion import generar_periodos
from apps.matriz.models import Area, EntidadControl, EventoPeriodo, Obligacion, Periodo
from apps.users.models import User

pytestmark = pytest.mark.django_db

#: Los roles que estas pruebas llamaban por su papel en la Matriz, con su
#: nombre en el catálogo general de roles.
ROL_GENERAL = {
    "Administrador": "Administrador",
    "Auditor": "Consulta / Auditoría",
    "Responsable": "Responsable de cumplimiento",
    "Supervisor/Aprobador": "Supervisor de cumplimiento",
}

GYE = zoneinfo.ZoneInfo("America/Guayaquil")


def _cuenta(nombre):
    return User.objects.create_user(
        username=nombre,
        email=f"{nombre}@example.com",
        password="Sup3r-Secr3t!",
        first_name=nombre.capitalize(),
    )


def _rol(nombre):
    return Group.objects.get(name=ROL_GENERAL[nombre])


def _miembro(usuario, empresa, rol):
    membresia = MembresiaEmpresa.objects.create(usuario=usuario, empresa=empresa)
    membresia.roles.set([_rol(rol)])
    return membresia


def _cliente(usuario, empresa):
    cliente = APIClient()
    cliente.force_authenticate(user=usuario)
    cliente.credentials(HTTP_X_EMPRESA=str(empresa.id))
    return cliente


@pytest.fixture
def courier():
    return Empresa.objects.create(nombre="LaarCourier", codigo="LC")


@pytest.fixture
def escena(courier):
    gente = {n: _cuenta(n) for n in ("admin", "ana", "beto", "sara", "aldo")}
    _miembro(gente["admin"], courier, "Administrador")
    _miembro(gente["ana"], courier, "Responsable")
    _miembro(gente["beto"], courier, "Responsable")
    _miembro(gente["sara"], courier, "Supervisor/Aprobador")
    _miembro(gente["aldo"], courier, "Auditor")
    clientes = {n: _cliente(u, courier) for n, u in gente.items()}
    return {"gente": gente, "empresa": courier, **clientes}


def _alta(escena, nombre, vence, responsable="ana", area="CONT", entidad="SRI", **extra):
    respuesta = escena["admin"].post(
        "/api/v1/matriz/obligaciones/",
        {
            "nombre": nombre,
            "area": Area.objects.get(codigo=area).id,
            "entidad": EntidadControl.objects.get(codigo=entidad).id,
            "periodicidad": extra.pop("periodicidad", "mensual"),
            "evidencia_esperada": "Formulario",
            "fecha_vencimiento": vence.isoformat(),
            "responsable": escena["gente"][responsable].id,
            **extra,
        },
        format="json",
    )
    assert respuesta.status_code == 201, respuesta.data
    return Periodo.objects.todas().get(pk=respuesta.data["id"])


def _cerrar(periodo, cumplido):
    Periodo.objects.todas().filter(pk=periodo.pk).update(
        cerrado=True,
        cumplido_en=cumplido,
        validado_en=cumplido,
        recordatorios_suspendidos=True,
    )


def _hoy():
    return timezone.now().astimezone(GYE).date()


# --- Generación de períodos (§11.3) ------------------------------------------------------


def test_genera_el_siguiente_y_hereda_del_ultimo(escena):
    primero = _alta(
        escena,
        "Declaración de IVA",
        _hoy() + datetime.timedelta(days=10),
        suplente=escena["gente"]["beto"].id,
        supervisor=escena["gente"]["sara"].id,
        prioridad="alta",
    )
    assert generar_periodos() == 1

    nuevo = Periodo.objects.todas().exclude(pk=primero.pk).get()
    assert nuevo.etiqueta != primero.etiqueta
    assert nuevo.vence_en.astimezone(GYE).time() == datetime.time(17, 0)
    assert (nuevo.responsable_id, nuevo.suplente_id, nuevo.supervisor_id) == (
        primero.responsable_id,
        primero.suplente_id,
        primero.supervisor_id,
    )
    assert nuevo.prioridad == "alta"
    evento = EventoPeriodo.objects.todas().get(periodo=nuevo)
    assert evento.descripcion == "Registro del período generado"
    assert evento.actor is None


def test_generar_es_idempotente(escena):
    _alta(escena, "Declaración de IVA", _hoy() + datetime.timedelta(days=10))
    assert generar_periodos() == 1
    assert generar_periodos() == 0
    assert Periodo.objects.todas().count() == 2


def test_lejos_del_vencimiento_no_genera(escena):
    _alta(escena, "Declaración de IVA", _hoy() + datetime.timedelta(days=45))
    assert generar_periodos() == 0


@pytest.mark.parametrize(
    "periodicidad,meses", [("mensual", 1), ("trimestral", 3), ("semestral", 6), ("anual", 12)]
)
def test_cada_periodicidad_suma_sus_meses(escena, periodicidad, meses):
    vence = _hoy() + datetime.timedelta(days=5)
    primero = _alta(escena, f"Obligación {periodicidad}", vence, periodicidad=periodicidad)
    generar_periodos()
    nuevo = Periodo.objects.todas().exclude(pk=primero.pk).order_by("vence_en").first()
    dia = nuevo.vence_en.astimezone(GYE).date()
    total = vence.month - 1 + meses
    assert (dia.year, dia.month) == (vence.year + total // 12, total % 12 + 1)


def test_la_unica_y_la_inactiva_no_generan(escena):
    _alta(escena, "Constitución", _hoy() + datetime.timedelta(days=3), periodicidad="unica")
    inactiva = _alta(escena, "Patente", _hoy() + datetime.timedelta(days=3))
    Obligacion.objects.todas().filter(pk=inactiva.obligacion_id).update(activa=False)
    assert generar_periodos() == 0


def test_no_crea_mas_de_24_por_pasada(escena):
    _alta(escena, "Atrasada", _hoy() - datetime.timedelta(days=365 * 4))
    assert generar_periodos() == 24


def test_un_periodo_adelantado_no_bloquea_la_generacion(escena):
    """«Noviembre» movido a octubre: el siguiente calculado es otra vez
    «Noviembre»; se salta a «Diciembre» en vez de no generar nunca más."""
    primero = _alta(escena, "Adelantada", _hoy().replace(day=1) + datetime.timedelta(days=40))
    Periodo.objects.todas().filter(pk=primero.pk).update(
        vence_en=timezone.now() - datetime.timedelta(days=3)
    )
    assert generar_periodos() == 1
    nuevo = Periodo.objects.todas().exclude(pk=primero.pk).get()
    assert nuevo.etiqueta != primero.etiqueta
    assert generar_periodos() == 0


def test_el_ciclo_genera_antes_de_avisar(escena):
    from apps.matriz.recordatorios import ciclo

    _alta(escena, "Declaración de IVA", _hoy() + datetime.timedelta(days=10))
    assert ciclo()["periodos"] == 1


# --- Calendario (§15.6) -------------------------------------------------------------------


def test_el_calendario_trae_lo_que_vence_en_el_mes(escena):
    _alta(escena, "Octubre", datetime.date(2099, 10, 31))
    _alta(escena, "Noviembre", datetime.date(2099, 11, 1))
    _alta(escena, "De Beto", datetime.date(2099, 10, 5), responsable="beto")

    respuesta = escena["admin"].get("/api/v1/matriz/calendario/?year=2099&month=10")
    assert respuesta.status_code == 200
    assert [p["obligacion"]["nombre"] for p in respuesta.data["periods"]] == [
        "De Beto",
        "Octubre",
    ]
    # Sin `ver_todas`, solo lo propio (§8.6).
    propio = escena["ana"].get("/api/v1/matriz/calendario/?year=2099&month=10")
    assert [p["obligacion"]["nombre"] for p in propio.data["periods"]] == ["Octubre"]
    assert escena["ana"].get("/api/v1/matriz/calendario/?month=13").status_code == 400


# --- Reportes (§16) -----------------------------------------------------------------------


@pytest.fixture
def con_cierres(escena):
    a_tiempo = _alta(escena, "A tiempo", datetime.date(2099, 1, 20))
    _cerrar(a_tiempo, a_tiempo.vence_en - datetime.timedelta(days=1))
    tarde = _alta(escena, "Tarde", datetime.date(2099, 1, 21), area="TH", entidad="IESS")
    _cerrar(tarde, tarde.vence_en + datetime.timedelta(days=2))
    otra = _alta(escena, "Otra tarde", datetime.date(2099, 1, 22), responsable="beto")
    _cerrar(otra, otra.vence_en + datetime.timedelta(days=1))
    _alta(escena, "Abierta", datetime.date(2099, 1, 25))
    vencida = _alta(escena, "Vencida", datetime.date(2099, 1, 26))
    Periodo.objects.todas().filter(pk=vencida.pk).update(
        vence_en=timezone.now() - datetime.timedelta(days=3)
    )
    return escena


def test_el_reporte_calcula_los_porcentajes_y_agrupa(con_cierres):
    respuesta = con_cierres["admin"].get("/api/v1/matriz/reportes/")
    assert respuesta.status_code == 200
    kpis = respuesta.data["kpis"]
    assert (kpis["on_time_pct"], kpis["late_pct"]) == (33, 67)
    assert (kpis["overdue"], kpis["in_progress"], kpis["closed"]) == (1, 1, 3)
    nombres = [fila["name"] for fila in respuesta.data["by_area"]]
    assert nombres == sorted(nombres, key=str.lower)
    contabilidad = next(f for f in respuesta.data["by_area"] if f["total"] == 4)
    assert contabilidad["finalizado"] == 2


def test_el_reporte_respeta_el_rango_y_el_alcance(con_cierres):
    rango = con_cierres["admin"].get(
        "/api/v1/matriz/reportes/?due_from=2099-01-20&due_to=2099-01-21"
    )
    assert rango.data["kpis"]["closed"] == 2
    propio = con_cierres["ana"].get("/api/v1/matriz/reportes/")
    assert propio.data["kpis"]["total"] == 4


def test_sin_cerrados_el_porcentaje_es_cero(escena):
    respuesta = escena["admin"].get("/api/v1/matriz/reportes/")
    assert respuesta.data["kpis"]["on_time_pct"] == 0


def test_exportar_exige_permiso_y_queda_auditado(con_cierres):
    assert con_cierres["ana"].get("/api/v1/matriz/reportes/?export=xlsx").status_code == 403

    excel = con_cierres["aldo"].get("/api/v1/matriz/reportes/?export=xlsx")
    assert excel.status_code == 200
    assert excel["Content-Type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in excel["Content-Disposition"]
    pdf = con_cierres["aldo"].get("/api/v1/matriz/reportes/?export=pdf")
    assert pdf.content.startswith(b"%PDF")
    assert AuditLog.objects.filter(action="report.exported").count() == 2
    assert con_cierres["aldo"].get("/api/v1/matriz/reportes/?export=doc").status_code == 400


def test_la_matriz_exportada_tiene_las_filas_de_la_pantalla(con_cierres):
    filtro = "status=finalizado"
    pantalla = con_cierres["admin"].get(f"/api/v1/matriz/periodos/?{filtro}")
    archivo = con_cierres["admin"].get(f"/api/v1/matriz/periodos/?{filtro}&export=xlsx")
    assert archivo.status_code == 200
    hoja = load_workbook(io.BytesIO(archivo.content)).active
    codigos = [fila[0].value for fila in hoja.iter_rows(min_row=6) if fila[0].value]
    assert codigos == [p["codigo"] for p in pantalla.data["results"]]
    contexto = hoja["A3"].value
    assert "LaarCourier" in contexto and "exportado por Admin" in contexto
    assert "estado: finalizado" in contexto


# --- Auditoría de la empresa (§14.2) -------------------------------------------------------


def test_la_auditoria_exige_ver_auditoria_y_filtra(escena):
    _alta(escena, "Declaración de IVA", datetime.date(2099, 1, 20))
    assert escena["ana"].get("/api/v1/matriz/auditoria/").status_code == 403

    respuesta = escena["aldo"].get("/api/v1/matriz/auditoria/?action=period.created")
    assert respuesta.status_code == 200
    assert respuesta.data["count"] == 1
    fila = respuesta.data["results"][0]
    assert fila["periodo"]["obligacion"] == "Declaración de IVA"
    assert fila["actor"] == "Admin"
    vacia = escena["aldo"].get("/api/v1/matriz/auditoria/?actor=sistema")
    assert vacia.data["count"] == 0


def test_la_auditoria_de_otra_empresa_no_se_ve(escena):
    _alta(escena, "Declaración de IVA", datetime.date(2099, 1, 20))
    otra = Empresa.objects.create(nombre="Otra", codigo="OT")
    intruso = _cuenta("intruso")
    _miembro(intruso, otra, "Auditor")
    respuesta = _cliente(intruso, otra).get("/api/v1/matriz/auditoria/")
    assert respuesta.data["count"] == 0


# --- Membresías (§8.7, §8.9) --------------------------------------------------------------


def _asignacion(empresa, rol, **extra):
    return [{"empresa_id": empresa.id, "roles": [_rol(rol).id], **extra}]


def test_guardar_la_ficha_conserva_las_areas_de_la_matriz(escena, courier):
    ana = escena["gente"]["ana"]
    contabilidad = Area.objects.get(codigo="CONT")
    asignar_empresas(
        actor=None,
        usuario=ana,
        asignaciones=_asignacion(courier, "Responsable", areas=[contabilidad.id]),
    )
    asignar_empresas(actor=None, usuario=ana, asignaciones=_asignacion(courier, "Responsable"))
    membresia = MembresiaEmpresa.objects.get(usuario=ana, empresa=courier)
    assert list(membresia.areas_matriz.all()) == [contabilidad]

    asignar_empresas(
        actor=None, usuario=ana, asignaciones=_asignacion(courier, "Responsable", areas=[])
    )
    assert not MembresiaEmpresa.objects.get(usuario=ana, empresa=courier).areas_matriz.exists()


def test_las_areas_viajan_por_la_api_de_usuarios(escena, courier):
    raiz = User.objects.create_superuser("raiz", "raiz@example.com", "Sup3r-Secr3t!")
    cliente = APIClient()
    cliente.force_authenticate(user=raiz)
    contabilidad = Area.objects.get(codigo="CONT")
    respuesta = cliente.post(
        f"/api/v1/admin/users/{escena['gente']['ana'].id}/empresas/",
        {"empresas": _asignacion(courier, "Responsable", areas=[contabilidad.id])},
        format="json",
    )
    assert respuesta.status_code == 200, respuesta.data
    assert respuesta.data["empresas"][0]["areas"] == [
        {"id": contabilidad.id, "nombre": contabilidad.nombre}
    ]


def test_una_empresa_no_se_queda_sin_quien_gestione_sus_miembros(escena, courier):
    admin = escena["gente"]["admin"]
    with pytest.raises(PermissionError, match="sin nadie que gestione"):
        asignar_empresas(actor=None, usuario=admin, asignaciones=_asignacion(courier, "Auditor"))
    # Nada cambió: la transacción se deshizo.
    assert MembresiaEmpresa.objects.get(usuario=admin, empresa=courier).roles.get() == _rol(
        "Administrador"
    )

    _miembro(_cuenta("otro"), courier, "Administrador")
    asignar_empresas(actor=None, usuario=admin, asignaciones=_asignacion(courier, "Auditor"))


def test_una_empresa_sin_gestores_no_bloquea_a_nadie(courier):
    persona = _cuenta("persona")
    _miembro(persona, courier, "Responsable")
    asignar_empresas(actor=None, usuario=persona, asignaciones=[])
    assert not MembresiaEmpresa.objects.filter(usuario=persona).exists()
