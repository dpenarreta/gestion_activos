"""Importar la matriz legal desde el Excel de Administración y la ficha de la
obligación (27/10/2026)."""

import datetime
import io
from decimal import Decimal

import openpyxl
import pytest
from django.contrib.auth.models import Group
from django.core.files.uploadedfile import SimpleUploadedFile
from rest_framework.test import APIClient

from apps.core.models import AuditLog
from apps.empresas.contexto import usando_empresa
from apps.empresas.models import Empresa, MembresiaEmpresa
from apps.matriz import importacion
from apps.matriz.models import (
    Area,
    EntidadControl,
    FichaObligacion,
    Obligacion,
    Periodo,
)
from apps.organizacion.models import Sede
from apps.users.models import User

pytestmark = pytest.mark.django_db

URL = "/api/v1/matriz/importar/"

CABECERA = [
    "Área",
    "Concepto",
    "Recurrencia",
    "Regional",
    "Entidad de Control",
    "Fecha de Emisión",
    "Fecha de Vencimiento",
    "Responsable",
    "Correo electrónico",
    "Dias de vigencia",
    "Tipo de Seguro",
    "Propietario",
    "Marca",
    "Modelo",
    "Tipo",
    "Placa",
    "Broker",
    "Aseguradora",
    "Vigencia Desde",
    "Vigencia Hasta",
    "Prima",
    "Valor Asegurado",
    "Empresa",
    "Polizas Conjuntas",
    "Detalle",
    "GARANTÍA",
    "CANON",
    "INCREMENTO",
]
CORREOS = "dani@grupolaar.com; andres@grupolaar.com"


def _fila(**valores):
    fila = dict.fromkeys(CABECERA, "N/A")
    fila.update({"Responsable": "Daniela Albán; Andres Lima", "Correo electrónico": CORREOS})
    fila.update(valores)
    return [fila[c] for c in CABECERA]


def _excel():
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "CONSOLIDADO"
    hoja.append(CABECERA)
    hoja.append(
        _fila(
            **{
                "Área": "Administración",
                "Concepto": "Patente",
                "Recurrencia": "Anual",
                "Regional": "AMBATO BODEGA",
                "Entidad de Control": "GAD Municipal Ambato",
                "Fecha de Emisión": datetime.datetime(2026, 1, 10),
                "Fecha de Vencimiento": datetime.datetime(2027, 1, 10),
            }
        )
    )
    # Sin fecha: se carga al 31/12/2026, por confirmar.
    hoja.append(
        _fila(
            **{
                "Área": "Administración",
                "Concepto": "Permiso de Bomberos",
                "Recurrencia": "Bianual",
                "Regional": "QUITO MATRIZ",
                "Entidad de Control": "GAD del Districto Metropolitano de Quito",
            }
        )
    )
    hoja.append(
        _fila(
            **{
                "Área": "Administración",
                "Concepto": "VEHICULO",
                "Recurrencia": "Anual",
                "Regional": "QUITO MATRIZ",
                "Tipo de Seguro": "VEHICULO",
                "Marca": "CHEVROLET",
                "Modelo": "DMAX 4X2",
                "Tipo": "CAMIONETA",
                "Placa": "pcx-8344",
                "Broker": "Multiapoyo Cía. Ltda",
                "Aseguradora": "Aseguradora del Sur C.A.",
                "Vigencia Desde": "01-06-2026",
                "Vigencia Hasta": "01-07-2027",
                "Prima": "2.80 por persona",
                "Valor Asegurado": 18000,
                "Empresa": "LAARCOURIER",
                "Polizas Conjuntas": 204662,
            }
        )
    )
    hoja.append(
        _fila(
            **{
                "Área": "Administración",
                "Concepto": "ARRIENDO",
                "Recurrencia": "Bianual",
                "Regional": "CUENCA SUCURSAL",
                "Fecha de Emisión": "2024-15-05",
                "Fecha de Vencimiento": "2025-15-04",
                "GARANTÍA": 700,
                "CANON": 388.89,
                "INCREMENTO": "Penduente negociacion de canon",
            }
        )
    )
    hoja.append(
        _fila(
            **{
                "Área": "Legal",
                "Concepto": "Nombramiento Gerente General",
                "Recurrencia": "Bianual",
                "Regional": "QUITO MATRIZ",
                "Entidad de Control": "Registro Mercantil",
                "Fecha de Vencimiento": datetime.datetime(2027, 2, 11),
            }
        )
    )
    arriendos = libro.create_sheet("ARRIENDOS")
    arriendos.append(["MATRIZ LEGAL Y DE CUMPLIMIENTO 2026"])
    arriendos.append([])
    arriendos.append(["ID", "DEPARTAMENTO", "REGIONAL", "VIGENCIA", "RENOVADO", "OBSERVACIONES"])
    arriendos.append(
        [1, "ADMINISTRACION", "CUENCA SUCURSAL", "2 años", "Sí", "Renovaion automatica"]
    )
    contabilidad = libro.create_sheet("CONTABILIDAD")
    contabilidad.append(["MATRIZ LEGAL Y DE CUMPLIMIENTO 2026"])
    contabilidad.append([])
    contabilidad.append(
        [
            "ID",
            "DEPARTAMENTO",
            "DOCUMENTO",
            "FECHA DE EMISIÓN",
            "FECHA DE CADUCIDAD",
            "MES DE RENOVACIÓN",
        ]
    )
    contabilidad.append(
        [
            1,
            "CONTABILIDAD",
            "Declaración IVA Mensual",
            datetime.datetime(2026, 7, 11),
            datetime.datetime(2026, 8, 11),
            "Mensual",
        ]
    )
    contabilidad.append(
        [
            2,
            "CONTABILIDAD",
            "Super Cias",
            datetime.datetime(2025, 9, 1),
            datetime.datetime(2026, 9, 1),
            "Anual",
        ]
    )
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _archivo():
    return SimpleUploadedFile(
        "MATRIZ LEGAL FINAL.xlsx",
        _excel(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@pytest.fixture
def courier():
    empresa = Empresa.objects.create(nombre="LaarCourier", codigo="LC")
    for codigo, nombre in [
        ("CONT", "Contabilidad y Tributario"),
        ("LEG", "Legal y Permisos"),
        ("SST", "Seguridad y Salud en el Trabajo"),
        ("SEG", "Seguros y Pólizas"),
        ("INF", "Arriendos e Infraestructura"),
        ("AMB", "Ambiental"),
    ]:
        Area.objects.update_or_create(codigo=codigo, defaults={"nombre": nombre})
    EntidadControl.objects.update_or_create(
        codigo="GDDMDQ", defaults={"nombre": "GAD del Districto Metropolitano de Quito"}
    )
    EntidadControl.objects.update_or_create(codigo="RM", defaults={"nombre": "Registro Mercantil"})
    Sede.objects.create(empresa=empresa, nombre="Matriz Quito", ciudad="Quito")
    return empresa


def _miembro(usuario, empresa, rol):
    membresia = MembresiaEmpresa.objects.create(usuario=usuario, empresa=empresa)
    membresia.roles.set([Group.objects.get(name=rol)])
    return usuario


@pytest.fixture
def gente(courier):
    admin = User.objects.create_user(username="admin", email="admin@grupolaar.com", password="x")
    _miembro(admin, courier, "Administrador")
    dani = User.objects.create_user(
        username="dani", email="dani@grupolaar.com", password="x", first_name="Daniela"
    )
    andres = User.objects.create_user(
        username="andres", email="andres@grupolaar.com", password="x", first_name="Andrés"
    )
    _miembro(dani, courier, "Responsable de cumplimiento")
    _miembro(andres, courier, "Responsable de cumplimiento")
    return {"admin": admin, "dani": dani, "andres": andres}


def _cliente(usuario, empresa):
    cliente = APIClient()
    cliente.force_authenticate(user=usuario)
    cliente.credentials(HTTP_X_EMPRESA=str(empresa.id))
    return cliente


def _obligaciones(empresa):
    with usando_empresa(empresa):
        return {o.nombre: o for o in Obligacion.objects.select_related("area", "entidad")}


def test_simular_no_guarda_y_dice_que_haria(courier, gente):
    respuesta = _cliente(gente["admin"], courier).post(URL, {"archivo": _archivo()})
    assert respuesta.status_code == 200, respuesta.data
    datos = respuesta.data
    assert datos["aplicado"] is None
    assert datos["filas"] == 7 and datos["crear"] == 7 and datos["errores"] == []
    assert datos["responsable"]["correo"] == "dani@grupolaar.com"
    assert datos["suplente"]["correo"] == "andres@grupolaar.com"
    # «CUENCA SUCURSAL» se crea con el nombre del catálogo («Sucursal Cuenca»).
    assert datos["sedes_nuevas"] == ["Ambato Bodega", "Sucursal Cuenca"]
    assert any("invertidos" in a for a in datos["avisos"])
    nombres = {o["nombre"] for o in datos["obligaciones"]}
    assert {
        "Patente · Ambato Bodega",
        "Permiso de Bomberos · Quito Matriz",
        "Seguro de vehículo · PCX-8344",
        "Contrato de arriendo · Cuenca Sucursal",
        "Nombramiento Gerente General",
        "Declaración IVA Mensual",
        "Super Cias",
    } == nombres
    assert _obligaciones(courier) == {}


def test_aplicar_crea_obligaciones_periodos_y_fichas(courier, gente):
    respuesta = _cliente(gente["admin"], courier).post(
        URL, {"archivo": _archivo(), "aplicar": "true"}
    )
    assert respuesta.status_code == 200, respuesta.data
    assert respuesta.data["aplicado"]["creadas"] == 7
    obligaciones = _obligaciones(courier)

    patente = obligaciones["Patente · Ambato Bodega"]
    assert (patente.area.codigo, patente.entidad.nombre, patente.periodicidad) == (
        "CONT",
        "GAD Municipal Ambato",
        "anual",
    )
    assert (patente.dia_vencimiento, patente.mes_vencimiento) == (10, 1)
    with usando_empresa(courier):
        periodo = Periodo.objects.get(obligacion=patente)
        assert periodo.responsable == gente["dani"] and periodo.suplente == gente["andres"]
        assert periodo.sede.nombre == "Ambato Bodega"
        bomberos = Periodo.objects.get(
            obligacion=obligaciones["Permiso de Bomberos · Quito Matriz"]
        )
    # Sin fecha: 31/12/2026, por confirmar; la entidad con la falta corregida
    # es la misma que ya estaba.
    assert bomberos.etiqueta == "Por confirmar"
    assert "por confirmar" in bomberos.observaciones
    assert bomberos.sede.nombre == "Matriz Quito"
    assert obligaciones["Permiso de Bomberos · Quito Matriz"].entidad.codigo == "GDDMDQ"
    assert obligaciones["Permiso de Bomberos · Quito Matriz"].periodicidad == "bienal"

    vehiculo = FichaObligacion.objects.todas().get(
        obligacion=obligaciones["Seguro de vehículo · PCX-8344"]
    )
    assert vehiculo.tipo == "seguro" and vehiculo.placa == "PCX-8344"
    assert vehiculo.aseguradora == "Aseguradora del Sur C.A."
    assert vehiculo.prima is None and vehiculo.prima_detalle == "2.80 por persona"
    assert vehiculo.valor_asegurado == Decimal("18000.00")
    assert vehiculo.vigencia_hasta == datetime.date(2027, 7, 1)
    assert obligaciones["Seguro de vehículo · PCX-8344"].area.codigo == "SEG"

    arriendo = obligaciones["Contrato de arriendo · Cuenca Sucursal"]
    ficha = FichaObligacion.objects.todas().get(obligacion=arriendo)
    assert (ficha.canon, ficha.garantia, ficha.duracion, ficha.renovado) == (
        Decimal("388.89"),
        Decimal("700.00"),
        "2 años",
        True,
    )
    assert ficha.incremento == "Pendiente negociación de canon"
    assert ficha.emitido_en == datetime.date(2024, 5, 15)
    assert Sede.objects.todas().filter(empresa=courier, nombre="Sucursal Cuenca").exists() is True

    assert obligaciones["Super Cias"].entidad.nombre == "Superintendencia de Compañías"
    assert obligaciones["Declaración IVA Mensual"].periodicidad == "mensual"
    assert AuditLog.objects.filter(action="matriz.imported").exists()
    assert AuditLog.objects.filter(action="obligation.created").count() == 7


def test_repetir_la_importacion_actualiza_y_no_duplica(courier, gente):
    cliente = _cliente(gente["admin"], courier)
    cliente.post(URL, {"archivo": _archivo(), "aplicar": "true"})
    segunda = cliente.post(URL, {"archivo": _archivo(), "aplicar": "true"})
    assert segunda.status_code == 200, segunda.data
    assert segunda.data["aplicado"] == {
        "creadas": 0,
        "actualizadas": 7,
        "temas_nuevos": 0,
        "entidades_nuevas": 0,
        "sedes_nuevas": 0,
    }
    assert len(_obligaciones(courier)) == 7


def test_reconoce_lo_cargado_antes_con_faltas_de_ortografia(courier, gente):
    with usando_empresa(courier):
        Obligacion.objects.create(
            empresa=courier,
            nombre="Contrato de arriendo · Cuenca Sucursal",
            area=Area.objects.get(codigo="INF"),
            entidad=EntidadControl.objects.get(codigo="RM"),
            periodicidad="anual",
            evidencia_esperada="Contrato",
        )
    with usando_empresa(courier):
        filas, _ = importacion.leer(io.BytesIO(_excel()))
        plan = importacion.planificar(filas, courier)
    acciones = {p.nombre: p.existente is not None for p in plan.propuestas}
    assert acciones["Contrato de arriendo · Cuenca Sucursal"] is True
    assert acciones["Patente · Ambato Bodega"] is False


def test_al_actualizar_completa_la_sucursal_que_faltaba(courier, gente):
    cliente = _cliente(gente["admin"], courier)
    cliente.post(URL, {"archivo": _archivo(), "aplicar": "true"})
    patente = _obligaciones(courier)["Patente · Ambato Bodega"]
    with usando_empresa(courier):
        Periodo.objects.filter(obligacion=patente).update(sede=None)
    cliente.post(URL, {"archivo": _archivo(), "aplicar": "true"})
    with usando_empresa(courier):
        assert Periodo.objects.get(obligacion=patente).sede.nombre == "Ambato Bodega"


def test_sin_la_cuenta_del_responsable_no_se_importa(courier, gente):
    gente["dani"].status = User.Status.DISABLED
    gente["dani"].save()
    cliente = _cliente(gente["admin"], courier)
    respuesta = cliente.post(URL, {"archivo": _archivo(), "aplicar": "true"})
    assert respuesta.status_code == 400
    assert any("dani@grupolaar.com" in e for e in respuesta.data["errores"])
    assert _obligaciones(courier) == {}


def test_importar_pide_editar_los_catalogos(courier, gente):
    respuesta = _cliente(gente["dani"], courier).post(URL, {"archivo": _archivo()})
    assert respuesta.status_code == 403


def test_un_archivo_que_no_es_la_matriz_se_rechaza(courier, gente):
    libro = openpyxl.Workbook()
    salida = io.BytesIO()
    libro.save(salida)
    archivo = SimpleUploadedFile("otro.xlsx", salida.getvalue())
    respuesta = _cliente(gente["admin"], courier).post(URL, {"archivo": archivo})
    assert respuesta.status_code == 400
    assert "CONSOLIDADO" in str(respuesta.data)


def test_la_ficha_se_ve_con_la_obligacion_y_se_edita(courier, gente):
    cliente = _cliente(gente["admin"], courier)
    cliente.post(URL, {"archivo": _archivo(), "aplicar": "true"})
    patente = _obligaciones(courier)["Patente · Ambato Bodega"]
    detalle = cliente.get(f"/api/v1/matriz/obligaciones/{patente.id}/")
    assert detalle.data["ficha"]["emitido_en"] == "2026-01-10"
    respuesta = cliente.put(
        f"/api/v1/matriz/obligaciones/{patente.id}/ficha/",
        {"tipo": "general", "emitido_en": "2026-02-01", "notas": "Pagada en ventanilla"},
        format="json",
    )
    assert respuesta.status_code == 200, respuesta.data
    assert respuesta.data["notas"] == "Pagada en ventanilla"
    assert AuditLog.objects.filter(action="obligation.ficha_updated").exists()
    malo = cliente.put(
        f"/api/v1/matriz/obligaciones/{patente.id}/ficha/",
        {"vigencia_desde": "2026-05-01", "vigencia_hasta": "2026-01-01"},
        format="json",
    )
    assert malo.status_code == 400


def test_si_faltan_los_temas_se_crean(courier, gente):
    """En producción el catálogo de temas estaba vacío (27/10/2026)."""
    Area.objects.filter(codigo__in=["SEG", "INF"]).delete()
    cliente = _cliente(gente["admin"], courier)
    simulado = cliente.post(URL, {"archivo": _archivo()})
    assert simulado.data["errores"] == []
    assert simulado.data["temas_nuevos"] == [
        "SEG — Seguros y Pólizas",
        "INF — Arriendos e Infraestructura",
    ]
    aplicado = cliente.post(URL, {"archivo": _archivo(), "aplicar": "true"})
    assert aplicado.data["aplicado"]["temas_nuevos"] == 2
    assert Area.objects.get(codigo="SEG").nombre == "Seguros y Pólizas"


# --- Límites del archivo (auditoría de seguridad, 29/10/2026) ---------------------


def _zip_con_entrada_enorme():
    """Un .xlsx de pocos KB que se descomprime en más del máximo («zip bomb»)."""
    import zipfile

    salida = io.BytesIO()
    with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as zip_:
        zip_.writestr("[Content_Types].xml", "<Types/>")
        zip_.writestr("xl/worksheets/sheet1.xml", b"0" * (importacion.MAX_BYTES_DESCOMPRIMIDOS + 1))
    return salida.getvalue()


def test_un_xlsx_que_se_descomprime_de_mas_se_rechaza_sin_abrirlo(courier, gente, monkeypatch):
    contenido = _zip_con_entrada_enorme()
    assert len(contenido) < importacion.MAX_BYTES
    abierto = []
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *a, **k: abierto.append(1))
    archivo = SimpleUploadedFile("matriz.xlsx", contenido)
    respuesta = _cliente(gente["admin"], courier).post(URL, {"archivo": archivo})
    assert respuesta.status_code == 400
    assert "demasiado grande" in str(respuesta.data)
    assert abierto == []


def test_un_archivo_de_mas_de_10_mb_se_rechaza(courier, gente):
    archivo = SimpleUploadedFile("matriz.xlsx", b"PK\x03\x04" + b"0" * importacion.MAX_BYTES)
    respuesta = _cliente(gente["admin"], courier).post(URL, {"archivo": archivo})
    assert respuesta.status_code == 400
    assert "10 MB" in str(respuesta.data)


def test_un_archivo_que_no_es_zip_se_rechaza(courier, gente):
    archivo = SimpleUploadedFile("matriz.xlsx", b"<html>no soy un excel</html>")
    respuesta = _cliente(gente["admin"], courier).post(URL, {"archivo": archivo})
    assert respuesta.status_code == 400
    assert ".xlsx válido" in str(respuesta.data)


def test_una_hoja_con_demasiadas_filas_se_rechaza(courier, gente, monkeypatch):
    monkeypatch.setattr(importacion, "MAX_FILAS", 3)
    respuesta = _cliente(gente["admin"], courier).post(URL, {"archivo": _archivo()})
    assert respuesta.status_code == 400
    assert "más de 3 filas" in str(respuesta.data)
