"""Rol de pagos → Roles sin abrir (26/10/2026).

Talento Humano (`roles_pago.seguimiento`) ve, de las personas habilitadas,
quiénes no han abierto su rol de un mes, con cómo contactarlas: el correo
institucional y el teléfono solo si autorizaron que las contacten por él.
Abrirlo es haberlo consultado ella misma, con su código; que Talento Humano
lo consulte por su código no cuenta. Las no habilitadas no aparecen.
"""

import io

from django.utils import timezone

from apps.empresas.contexto import empresa_actual, usando_empresa
from apps.nomina.models import ConsultaRolPagos
from apps.organizacion.models import DatosDeContacto
from apps.organizacion.tests.test_carga_personas import _cuenta

URL = "/api/v1/nomina/rol-pagos/seguimiento/"


def _habilitada(nombre, codigo, *permisos):
    usuario, cliente = _cuenta(nombre, "roles_pago.consultar", *permisos)
    usuario.codigo_empleado = codigo
    usuario.save()
    return usuario, cliente


def _consulta(usuario, codigo, anio=2026, mes=9, **extra):
    with usando_empresa(empresa_actual()):
        return ConsultaRolPagos.objects.create(
            usuario=usuario,
            codigo_empleado=codigo,
            anio=anio,
            mes=mes,
            version_terminos=1,
            terminos_aceptados="…",
            **extra,
        )


def test_lista_quienes_no_abrieron_su_rol_y_como_contactarlas(db):
    ana, _ = _habilitada("ana", "100")
    beto, _ = _habilitada("beto", "200")
    carla, _ = _habilitada("carla", "300")
    _cuenta("dario")  # no habilitado: no aparece
    th, cliente = _habilitada("th", "900", "roles_pago.seguimiento", "roles_pago.consultar_todos")
    with usando_empresa(empresa_actual()):
        DatosDeContacto.objects.create(
            usuario=beto, telefono="0999999999", telefono_aceptado_en=timezone.now()
        )
        DatosDeContacto.objects.create(usuario=carla, telefono="0988888888")  # sin autorizar

    _consulta(ana, "100", conformidad="conforme", conformidad_en=timezone.now())
    _consulta(th, "200")  # Talento Humano abrió el de Beto: no cuenta
    _consulta(beto, "200", mes=8)  # el del mes anterior

    datos = cliente.get(URL, {"anio": 2026, "mes": 9}).json()

    assert datos["periodo"] == "Septiembre 2026"
    assert datos["habilitados"] == 4
    sin_abrir = {f["nombre"]: f for f in datos["sin_abrir"]}
    assert set(sin_abrir) == {"Beto Prueba", "Carla Prueba", "Th Prueba"}
    assert sin_abrir["Beto Prueba"]["telefono"] == "0999999999"
    assert sin_abrir["Beto Prueba"]["correo"] == "beto@grupolaar.com"
    assert sin_abrir["Beto Prueba"]["ultima_apertura"] is not None
    assert sin_abrir["Carla Prueba"]["telefono"] == ""
    assert sin_abrir["Carla Prueba"]["ultima_apertura"] is None
    (abierta,) = datos["aceptados"]
    assert datos["totales"] == {
        "sin_abrir": 3,
        "sin_aceptar": 0,
        "con_observacion": 0,
        "aceptados": 1,
    }
    assert abierta["nombre"] == "Ana Prueba" and abierta["codigo"] == "100"
    assert abierta["conformidad_display"] == "Estuvo de acuerdo"
    assert abierta["abierto"]["fecha"]


def test_sin_el_permiso_no_se_ve_y_el_mes_futuro_se_niega(db):
    _, persona = _habilitada("ana", "100")
    assert persona.get(URL, {"anio": 2026, "mes": 9}).status_code == 403
    _, th = _cuenta("th", "roles_pago.seguimiento")
    assert th.get(URL, {"anio": 2026, "mes": "x"}).status_code == 400
    assert th.get(URL, {"anio": 2999, "mes": 1}).status_code == 400


def test_la_pantalla_dice_si_puede_hacer_seguimiento(db):
    _, th = _cuenta("th", "roles_pago.seguimiento")
    estado = th.get("/api/v1/nomina/rol-pagos/")
    assert estado.status_code == 200
    assert estado.json()["puede_seguimiento"] is True


def test_la_persona_ve_cuando_abrio_su_rol_y_cuando_dio_su_conformidad(db):
    """Abrirlo es aceptar los términos: queda con fecha, hora e IP; la
    conformidad, con las suyas. Lo que consultó con otro código no es suyo."""
    ana, cliente = _habilitada("ana", "100", "roles_pago.consultar_todos")
    _consulta(ana, "100", ip="10.0.0.5")
    _consulta(
        ana,
        "100",
        ip="10.0.0.6",
        conformidad="conforme",
        conformidad_en=timezone.now(),
        conformidad_ip="10.0.0.7",
    )
    _consulta(ana, "100", mes=8, ip="10.0.0.5")
    _consulta(ana, "999", mes=7)  # el de otra persona

    mis = cliente.get("/api/v1/nomina/rol-pagos/").json()["mis_roles"]

    assert [(r["anio"], r["mes"]) for r in mis] == [(2026, 9), (2026, 8)]
    septiembre, agosto = mis
    assert septiembre["abierto"]["ip"] == "10.0.0.5"  # la primera vez
    assert septiembre["abierto"]["fecha"] and septiembre["abierto"]["hora"]
    assert septiembre["veces"] == 2
    assert septiembre["conformidad_display"] == "Estuvo de acuerdo"
    assert septiembre["conformidad_en"]["ip"] == "10.0.0.7"
    assert agosto["conformidad"] == "" and agosto["conformidad_en"] is None


def test_talento_humano_ve_la_ip_de_apertura_y_de_conformidad(db):
    ana, _ = _habilitada("ana", "100")
    _, th = _cuenta("th", "roles_pago.seguimiento")
    _consulta(
        ana,
        "100",
        ip="10.0.0.5",
        conformidad="observacion",
        conformidad_en=timezone.now(),
        conformidad_ip="10.0.0.9",
    )
    (fila,) = th.get(URL, {"anio": 2026, "mes": 9}).json()["con_observacion"]
    assert fila["abierto"]["ip"] == "10.0.0.5"
    assert fila["conformidad_en"]["ip"] == "10.0.0.9"
    assert fila["conformidad_display"] == "Dejó una observación"


def _mes_variado():
    """Septiembre: Ana aceptó, Beto lo abrió sin responder, Carla observó y
    Dario no lo abrió."""
    ana, _ = _habilitada("ana", "100")
    beto, _ = _habilitada("beto", "200")
    carla, _ = _habilitada("carla", "300")
    dario, _ = _habilitada("dario", "400")
    _, th = _cuenta("th", "roles_pago.seguimiento")
    ahora = timezone.now()
    _consulta(ana, "100", conformidad="conforme", conformidad_en=ahora)
    _consulta(beto, "200")
    _consulta(carla, "300", conformidad="observacion", conformidad_en=ahora)
    _consulta(ana, "100", mes=8, conformidad="conforme", conformidad_en=ahora)
    return th, dario


def test_cada_mes_separa_sin_abrir_sin_aceptar_observados_y_aceptados(db):
    th, _ = _mes_variado()
    datos = th.get(URL, {"anio": 2026, "mes": 9}).json()
    assert [f["nombre"] for f in datos["sin_abrir"]] == ["Dario Prueba"]
    assert [f["nombre"] for f in datos["sin_aceptar"]] == ["Beto Prueba"]
    assert [f["nombre"] for f in datos["con_observacion"]] == ["Carla Prueba"]
    assert [f["nombre"] for f in datos["aceptados"]] == ["Ana Prueba"]


def test_el_resumen_cuenta_por_mes(db):
    th, _ = _mes_variado()
    filas = th.get(f"{URL}resumen/", {"anio": 2026}).json()
    por_mes = {f["mes"]: f for f in filas}
    assert por_mes[9]["habilitados"] == 4
    assert por_mes[9]["abiertos"] == 3 and por_mes[9]["aceptados"] == 1
    assert por_mes[9]["sin_abrir"] == 1 and por_mes[9]["pendientes"] == 2
    assert por_mes[8]["aceptados"] == 1 and por_mes[8]["sin_abrir"] == 3
    assert filas[0]["mes"] > filas[-1]["mes"]  # lo más reciente primero


def test_el_listado_se_exporta_a_excel(db):
    from openpyxl import load_workbook

    th, _ = _mes_variado()
    respuesta = th.get(f"{URL}exportar/", {"anio": 2026, "mes": 9})
    assert respuesta.status_code == 200
    assert "roles-2026-09-pendientes.xlsx" in respuesta["Content-Disposition"]
    hoja = load_workbook(io.BytesIO(respuesta.content)).active
    filas = list(hoja.iter_rows(min_row=5, values_only=True))
    assert [(f[0], f[1], f[2], f[3]) for f in filas] == [
        ("Sin abrir", "Dario Prueba", "400", "dario@grupolaar.com"),
        ("Abierto, sin aceptar", "Beto Prueba", "200", "beto@grupolaar.com"),
    ]
    todos = th.get(f"{URL}exportar/", {"anio": 2026, "mes": 9, "estado": "todos"})
    assert len(list(load_workbook(io.BytesIO(todos.content)).active.iter_rows(min_row=5))) == 4
    assert th.get(f"{URL}exportar/", {"anio": 2026, "mes": 9, "estado": "x"}).status_code == 400


def test_el_recordatorio_llega_a_los_pendientes_una_vez(db):
    from apps.notificaciones.models import Notificacion

    th, dario = _mes_variado()
    respuesta = th.post(f"{URL}recordar/", {"anio": 2026, "mes": 9}, format="json")
    assert respuesta.json() == {"enviadas": 2}
    aviso = Notificacion.objects.get(usuario=dario)
    assert "no ha abierto" in aviso.resumen
    assert aviso.enlace == "/admin/institucional/rol-pagos?anio=2026&mes=9"
    # Repetirlo enseguida no vuelve a avisar.
    assert th.post(f"{URL}recordar/", {"anio": 2026, "mes": 9}, format="json").json() == {
        "enviadas": 0
    }
    assert (
        th.post(
            f"{URL}recordar/", {"anio": 2026, "mes": 9, "estado": "aceptados"}, format="json"
        ).status_code
        == 400
    )
