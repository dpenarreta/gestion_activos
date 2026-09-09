"""Genera el .xlsx de un catálogo: instrucciones, columnas y lo que ya existe.

El archivo baja **lleno**. Es deliberado y hace dos trabajos con uno solo: sirve
de referencia mientras se llena la plantilla de activos —ahí es donde hay que
copiar el nombre exacto de una sede o el código de un área— y sirve de
plantilla, porque las filas nuevas se escriben debajo de las que ya están.

Que se pueda bajar, añadir tres filas y volver a subir sin borrar nada es la
razón de que el importador reconozca lo que ya existe en vez de tratarlo como
un duplicado.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .catalogos_masivos import CatalogoMasivo

RELLENO_CABECERA = PatternFill("solid", fgColor="1F3864")
RELLENO_OBLIGATORIA = PatternFill("solid", fgColor="7F1D1D")
FUENTE_CABECERA = Font(bold=True, color="FFFFFF")


def _hoja_datos(libro: Workbook, catalogo: CatalogoMasivo) -> None:
    hoja = libro.active
    hoja.title = catalogo.nombre

    for indice, columna in enumerate(catalogo.columnas, start=1):
        celda = hoja.cell(row=1, column=indice, value=columna.encabezado)
        celda.font = FUENTE_CABECERA
        celda.fill = RELLENO_OBLIGATORIA if columna.obligatoria else RELLENO_CABECERA
        celda.alignment = Alignment(horizontal="center", vertical="center")
        if columna.ayuda:
            # La ayuda va en el encabezado y no en una hoja aparte: se lee justo
            # cuando hace falta, con el cursor ya en la columna.
            celda.comment = Comment(columna.ayuda, "Gestión de Activos")
        hoja.column_dimensions[get_column_letter(indice)].width = columna.ancho

    Modelo = catalogo.obtener_modelo()
    existentes = Modelo.objects.filter(**catalogo.filtro_vigentes).order_by(*catalogo.orden)
    for numero_fila, objeto in enumerate(existentes, start=2):
        for indice, columna in enumerate(catalogo.columnas, start=1):
            hoja.cell(row=numero_fila, column=indice, value=_valor(objeto, columna))

    hoja.freeze_panes = "A2"


def _valor(objeto, columna) -> str:
    if columna.sensible:
        # La columna existe para llenarla, no para llevarse lo que ya hay: el
        # archivo se descarga y circula fuera del sistema.
        return ""
    valor = getattr(objeto, columna.clave, "")
    if valor is None:
        return ""
    if columna.referencia:
        # Se escribe como se lee, no con el identificador: es lo que hay que
        # teclear para volver a apuntar ahí.
        campo = columna.referencia.campos[0]
        return str(getattr(valor, campo, valor))
    return str(valor)


def _hoja_instrucciones(libro: Workbook, catalogo: CatalogoMasivo) -> None:
    hoja = libro.create_sheet("Instrucciones", 0)
    hoja.column_dimensions["A"].width = 108

    lineas = [
        (f"Carga masiva de {catalogo.plural}", True),
        ("", False),
        (
            f"1. La hoja «{catalogo.nombre}» baja con lo que ya está registrado. "
            "Escriba las filas nuevas debajo, sin borrar las que hay.",
            False,
        ),
        (
            "2. Volver a subir este mismo archivo no duplica nada: las filas que ya "
            "existen se reconocen y se omiten.",
            False,
        ),
        ("3. Las columnas marcadas con * son obligatorias.", False),
        (
            "4. Pase el cursor sobre cada encabezado para ver qué se espera en esa columna.",
            False,
        ),
        (
            "5. Nada se guarda al subir el archivo: primero se revisa entero y se le "
            "muestra qué entraría. Recién entonces se confirma.",
            False,
        ),
    ]
    if catalogo.nota:
        lineas.append((f"6. {catalogo.nota}", False))

    for numero_fila, (texto, es_titulo) in enumerate(lineas, start=1):
        celda = hoja.cell(row=numero_fila, column=1, value=texto)
        celda.font = Font(bold=True, size=14) if es_titulo else Font(size=11)
        celda.alignment = Alignment(wrap_text=True, vertical="top")


def construir_plantilla(catalogo: CatalogoMasivo) -> bytes:
    libro = Workbook()
    _hoja_datos(libro, catalogo)
    _hoja_instrucciones(libro, catalogo)

    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()
