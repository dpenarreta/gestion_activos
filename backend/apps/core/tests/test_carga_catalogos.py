"""Carga masiva de los catálogos: un archivo por catálogo.

La plantilla de activos llegó a traer ocho hojas en un solo libro, cinco de
ellas solo de consulta, y ninguno de esos catálogos se podía cargar: para dar
de alta cincuenta empleados había que teclearlos uno a uno mientras el archivo
ya los listaba.
"""

from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from apps.activos.models import TipoDispositivo
from apps.core.catalogos_masivos import POR_CLAVE
from apps.organizacion.models import Departamento, Empleado, Proveedor, Sede
from apps.users.models import User

BASE = "/api/v1/catalogos"


@pytest.fixture
def admin(db):
    return User.objects.create_superuser(
        username="admin_cat", email="admin_cat@example.com", password="Sup3r-Secr3t!"
    )


@pytest.fixture
def cliente(admin):
    client = APIClient()
    client.force_authenticate(user=admin)
    return client


def _archivo(catalogo, filas):
    """Un .xlsx con la forma que produce la plantilla."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = catalogo.nombre
    hoja.append([columna.encabezado for columna in catalogo.columnas])
    for fila in filas:
        hoja.append([fila.get(columna.clave, "") for columna in catalogo.columnas])
    buffer = BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    buffer.name = "carga.xlsx"
    return buffer


def _subir(cliente, clave, archivo, confirmar=False):
    return cliente.post(
        f"{BASE}/{clave}/importar/",
        {"archivo": archivo, "confirmar": "true" if confirmar else "false"},
        format="multipart",
    )


# --- Un archivo por catálogo ------------------------------------------------


def test_cada_catalogo_tiene_su_propio_archivo(cliente):
    """Y solo el suyo: material de trabajo separado del de consulta."""
    for clave, catalogo in POR_CLAVE.items():
        respuesta = cliente.get(f"{BASE}/{clave}/plantilla/")

        assert respuesta.status_code == 200, clave
        libro = load_workbook(BytesIO(respuesta.content))
        assert libro.sheetnames == ["Instrucciones", catalogo.nombre], clave


def test_la_plantilla_de_activos_ya_no_arrastra_los_catalogos(cliente, db):
    respuesta = cliente.get("/api/v1/activos/plantilla-importacion/")

    libro = load_workbook(BytesIO(respuesta.content))
    assert libro.sheetnames == ["Instrucciones", "Activos", "Ejemplo"]


def test_el_archivo_baja_con_lo_que_ya_existe_dentro(cliente, db):
    """Hace dos trabajos con uno: sirve de referencia al llenar la plantilla de
    activos y de plantilla para añadir filas debajo."""
    Sede.objects.create(nombre="Sede Quito Norte", ciudad="Quito")
    Sede.objects.create(nombre="Sede Guayaquil", ciudad="Guayaquil")

    respuesta = cliente.get(f"{BASE}/sedes/plantilla/")

    hoja = load_workbook(BytesIO(respuesta.content))["Sedes"]
    nombres = [hoja.cell(row=fila, column=1).value for fila in range(2, hoja.max_row + 1)]
    assert nombres == ["Sede Guayaquil", "Sede Quito Norte"]


# --- Cargar ----------------------------------------------------------------


def test_crea_las_filas_nuevas(cliente, db):
    archivo = _archivo(
        POR_CLAVE["sedes"],
        [
            {"nombre": "Sede Quito Norte", "ciudad": "Quito"},
            {"nombre": "Sede Guayaquil", "ciudad": "Guayaquil", "direccion": "Av. 9 de Octubre"},
        ],
    )

    respuesta = _subir(cliente, "sedes", archivo, confirmar=True)

    assert respuesta.status_code == 201
    assert respuesta.json()["creados"] == 2
    assert Sede.objects.get(nombre="Sede Guayaquil").ciudad == "Guayaquil"


def test_volver_a_subir_el_archivo_descargado_no_duplica_nada(cliente, db):
    """Es lo que la gente hace en cuanto descubre que el archivo se reutiliza:
    bajarlo, añadir tres filas y volver a subirlo entero."""
    Sede.objects.create(nombre="Sede Quito Norte", ciudad="Quito")
    archivo = _archivo(
        POR_CLAVE["sedes"],
        [
            {"nombre": "Sede Quito Norte", "ciudad": "Quito"},
            {"nombre": "Sede Cuenca", "ciudad": "Cuenca"},
        ],
    )

    reporte = _subir(cliente, "sedes", archivo, confirmar=True).json()

    assert reporte["creados"] == 1
    assert Sede.objects.count() == 2
    assert "Ya existe" in reporte["advertencias"][0]["mensaje"]


def test_nada_se_guarda_sin_confirmar(cliente, db):
    """Dos pasos, como en la carga de activos: importar directo dejaría al
    usuario descubriendo los errores con los registros ya creados."""
    archivo = _archivo(POR_CLAVE["sedes"], [{"nombre": "Sede Nueva", "ciudad": "Loja"}])

    reporte = _subir(cliente, "sedes", archivo).json()

    assert reporte["filas_validas"] == 1
    assert reporte["es_importable"] is True
    assert Sede.objects.count() == 0


# --- Lo que se rechaza -----------------------------------------------------


def test_una_columna_obligatoria_vacia_bloquea_la_fila(cliente, db):
    archivo = _archivo(POR_CLAVE["sedes"], [{"ciudad": "Quito"}])

    reporte = _subir(cliente, "sedes", archivo).json()

    assert reporte["es_importable"] is False
    assert reporte["errores"][0]["columna"] == "Nombre"


def test_el_repetido_dentro_del_archivo_si_bloquea(cliente, db):
    """A diferencia del que ya existe: dos filas iguales en el mismo archivo son
    un error de quien lo armó, no una carga repetida."""
    archivo = _archivo(
        POR_CLAVE["departamentos"],
        [
            {"codigo": "TI", "nombre": "Tecnología"},
            {"codigo": "ti", "nombre": "Sistemas"},
        ],
    )

    reporte = _subir(cliente, "departamentos", archivo).json()

    assert reporte["es_importable"] is False
    assert "fila 2" in reporte["errores"][0]["mensaje"]


def test_una_referencia_que_no_existe_dice_donde_crearla(cliente, db):
    archivo = _archivo(
        POR_CLAVE["empleados"],
        [{"nombres": "Ana", "apellidos": "Pérez", "departamento": "NO-EXISTE"}],
    )

    reporte = _subir(cliente, "empleados", archivo).json()

    assert reporte["es_importable"] is False
    assert "Cárguelo primero" in reporte["errores"][0]["mensaje"]


# --- Referencias y reglas del modelo ---------------------------------------


def test_el_area_se_escribe_por_codigo_o_por_nombre(cliente, db):
    """Quien llena la plantilla escribe lo que tiene en la cabeza."""
    Departamento.objects.create(nombre="Tecnología", codigo="TI")
    archivo = _archivo(
        POR_CLAVE["empleados"],
        [
            {"nombres": "Ana", "apellidos": "Pérez", "departamento": "TI"},
            {"nombres": "Luis", "apellidos": "Torres", "departamento": "Tecnología"},
        ],
    )

    assert _subir(cliente, "empleados", archivo, confirmar=True).json()["creados"] == 2


def test_el_codigo_del_empleado_se_genera_al_cargar(cliente, db):
    """Se crea uno a uno y no con `bulk_create` justo por esto: el código sale
    del guardado, y saltárselo dejaría filas sin código en la carga que más
    filas produce."""
    Departamento.objects.create(nombre="Tecnología", codigo="TI")
    archivo = _archivo(
        POR_CLAVE["empleados"],
        [
            {"nombres": "Ana", "apellidos": "Pérez", "departamento": "TI"},
            {"nombres": "Luis", "apellidos": "Torres", "departamento": "TI"},
        ],
    )

    _subir(cliente, "empleados", archivo, confirmar=True)

    assert sorted(Empleado.objects.values_list("codigo_empleado", flat=True)) == [
        "TI-0001",
        "TI-0002",
    ]


def test_la_carga_es_todo_o_nada(cliente, db):
    """Una carga a medias deja al usuario sin saber qué entró."""
    Departamento.objects.create(nombre="Tecnología", codigo="TI")
    archivo = _archivo(
        POR_CLAVE["empleados"],
        [
            {"nombres": "Ana", "apellidos": "Pérez", "departamento": "TI"},
            {"nombres": "Luis", "apellidos": "", "departamento": "TI"},
        ],
    )

    respuesta = _subir(cliente, "empleados", archivo, confirmar=True)

    assert respuesta.status_code == 400
    assert Empleado.objects.count() == 0


# --- Permisos ---------------------------------------------------------------


def test_cada_catalogo_pide_el_permiso_de_su_modulo(db):
    """Cargar tipos de dispositivo es editar el inventario; cargar empleados es
    editar la organización. Un permiso único de «carga masiva» daría acceso a
    los dos a quien solo necesita uno."""
    assert POR_CLAVE["tipos"].permiso == "activos.editar"
    assert POR_CLAVE["empleados"].permiso == "organizacion.editar"


def test_sin_permiso_no_se_descarga_ni_se_carga(db):
    usuario = User.objects.create_user(
        username="pelado_cat", email="pelado_cat@example.com", password="Sup3r-Secr3t!"
    )
    cliente = APIClient()
    cliente.force_authenticate(user=usuario)

    assert cliente.get(f"{BASE}/sedes/plantilla/").status_code == 403
    assert _subir(cliente, "sedes", _archivo(POR_CLAVE["sedes"], [])).status_code == 403


def test_un_catalogo_que_no_existe_devuelve_404(cliente):
    assert cliente.get(f"{BASE}/inventado/plantilla/").status_code == 404


def test_el_listado_dice_que_se_puede_cargar(cliente):
    """La pantalla se dibuja desde aquí: un catálogo nuevo aparece sin tocar el
    frontend."""
    datos = cliente.get(f"{BASE}/").data

    claves = {c["clave"] for c in datos["catalogos"]}
    assert claves == {"sedes", "departamentos", "proveedores", "tipos", "empleados"}
    assert all(c["puede"] for c in datos["catalogos"])


# --- Los otros catálogos ----------------------------------------------------


def test_se_cargan_proveedores_y_tipos(cliente, db):
    proveedores = _archivo(
        POR_CLAVE["proveedores"],
        [{"nombre": "Tecnomega", "identificacion": "0991234567001", "telefono": "04-2000000"}],
    )
    tipos = _archivo(POR_CLAVE["tipos"], [{"codigo": "LAP", "nombre": "Laptop"}])

    assert _subir(cliente, "proveedores", proveedores, confirmar=True).status_code == 201
    assert _subir(cliente, "tipos", tipos, confirmar=True).status_code == 201
    assert Proveedor.objects.get().telefono == "04-2000000"
    assert TipoDispositivo.objects.get().codigo == "LAP"


def test_el_archivo_de_otro_formato_se_rechaza_con_mensaje(cliente, db):
    from django.core.files.uploadedfile import SimpleUploadedFile

    respuesta = cliente.post(
        f"{BASE}/sedes/importar/",
        {"archivo": SimpleUploadedFile("datos.csv", b"nombre,ciudad")},
        format="multipart",
    )

    assert respuesta.status_code == 400
    assert respuesta.data["error"]["code"] == "formato_invalido"
